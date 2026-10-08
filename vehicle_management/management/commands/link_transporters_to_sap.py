"""Tie the app's transporters to their SAP vendors, from a reviewed mapping sheet.

The sheet is ``Transporter_SAP_mapping_<date>.xlsx``, sheet "Mapping": one row
per app transporter, with the suggested match and its Oil, Mart and Beverages
vendor codes. Which rows are linked:

  * "GSTIN match" and "Same name" rows, unless "Your call" says manual, no or skip;
  * any other row only when "Your call" says OK (fix the code cells first if the
    suggestion is wrong; a "Not in SAP" row needs its codes typed in).

Every code is checked against that company's SAP. A frozen vendor is still
linked: the link says who the transporter is, not that it can be picked. A blank
transporter GSTIN is filled from SAP; a different one is left alone and listed.

Dry run by default; pass --apply to write. Needs HANA.

    manage.py link_transporters_to_sap ~/Transporter_SAP_mapping_2026-10-08.xlsx --apply
"""

from collections import Counter

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from openpyxl import load_workbook

from company.models import Company
from vehicle_management import sap_transporters
from vehicle_management.models import Transporter, TransporterSAPLink

SHEET = "Mapping"
# Sheet column -> company code.
CODE_COLUMNS = {
    "Oil code": "JIVO_OIL",
    "Mart code": "JIVO_MART",
    "Beverages code": "JIVO_BEVERAGES",
}
SURE_MATCHES = {"GSTIN match", "Same name"}
YES = {"ok", "yes", "y"}
NO = {"manual", "no", "n", "skip", "x"}


def _code(cell):
    """``VENDA000335 (frozen)`` -> ``VENDA000335``."""
    return str(cell or "").replace("(frozen)", "").strip().upper()


class Command(BaseCommand):
    help = "Link app transporters to SAP vendors from a reviewed mapping sheet (dry run unless --apply)."

    def add_arguments(self, parser):
        parser.add_argument("sheet", help="Path to the reviewed Transporter_SAP_mapping xlsx.")
        parser.add_argument("--apply", action="store_true")

    def handle(self, *args, **options):
        apply_changes = options["apply"]
        self.stdout.write("APPLY" if apply_changes else "DRY RUN (nothing written)")

        plan, skipped, problems = self._read_sheet(options["sheet"])
        companies = {c.code: c for c in Company.objects.filter(code__in=CODE_COLUMNS.values())}
        missing = set(CODE_COLUMNS.values()) - set(companies)
        if missing:
            raise CommandError(f"No company {', '.join(sorted(missing))} in this database.")

        vendors = {}
        for company_code in sorted({company_code for _, company_code, _ in plan}):
            vendors[company_code] = sap_transporters.vendors_by_code(company_code)
            self.stdout.write(f"{company_code}: {len(vendors[company_code])} SAP vendors read")

        transporters = Transporter.objects.in_bulk({transporter_id for transporter_id, _, _ in plan})
        existing = set(
            TransporterSAPLink.objects.values_list("transporter_id", "company__code", "card_code")
        )

        counts = Counter()
        gstin_differs = []
        with transaction.atomic():
            for transporter_id, company_code, card_code in plan:
                transporter = transporters.get(transporter_id)
                if transporter is None:
                    problems.append(f"#{transporter_id}: no such transporter")
                    continue
                vendor = vendors[company_code].get(card_code)
                if vendor is None:
                    problems.append(
                        f"#{transporter_id} {transporter.name}: {card_code} is not a vendor in {company_code}"
                    )
                    continue
                if (transporter_id, company_code, card_code) in existing:
                    counts["already linked"] += 1
                    continue
                counts[f"links to make ({company_code})"] += 1
                if vendor["gstin"] and transporter.gstin and transporter.gstin.upper() != vendor["gstin"]:
                    gstin_differs.append(
                        f"#{transporter_id} {transporter.name}: app {transporter.gstin}, "
                        f"SAP {company_code} {card_code} {vendor['gstin']}"
                    )
                elif vendor["gstin"] and not transporter.gstin:
                    counts["GSTINs to fill"] += 1
                if apply_changes:
                    sap_transporters.link(transporter, companies[company_code], vendor)
                existing.add((transporter_id, company_code, card_code))

        for key in sorted(counts):
            self.stdout.write(f"{key}: {counts[key]}")
        for reason, n in sorted(skipped.items()):
            self.stdout.write(f"rows skipped, {reason}: {n}")
        if gstin_differs:
            self.stdout.write("App GSTIN differs from SAP's (left as it is):")
            for line in gstin_differs:
                self.stdout.write(f"  {line}")
        if problems:
            self.stdout.write(self.style.WARNING("Not linked:"))
            for line in problems:
                self.stdout.write(f"  {line}")
        self.stdout.write(self.style.SUCCESS("Done." if apply_changes else "Dry run done; --apply to write."))

    def _read_sheet(self, path):
        try:
            workbook = load_workbook(path, read_only=True, data_only=True)
        except Exception as exc:  # noqa: BLE001 -- any unreadable file is the same answer
            raise CommandError(f"Cannot read {path}: {exc}") from exc
        if SHEET not in workbook.sheetnames:
            raise CommandError(f"{path} has no '{SHEET}' sheet.")
        rows = workbook[SHEET].iter_rows(values_only=True)
        header = [str(cell or "").strip() for cell in next(rows)]
        needed = ["App ID", "Match", "Your call", *CODE_COLUMNS]
        absent = [name for name in needed if name not in header]
        if absent:
            raise CommandError(f"The '{SHEET}' sheet has no column {', '.join(absent)}.")
        at = {name: header.index(name) for name in needed}

        plan, skipped, problems = [], Counter(), []
        for row in rows:
            if row is None or row[at["App ID"]] in (None, ""):
                continue
            try:
                transporter_id = int(row[at["App ID"]])
            except (TypeError, ValueError):
                problems.append(f"App ID {row[at['App ID']]!r} is not a number")
                continue
            match = str(row[at["Match"]] or "").strip()
            call = str(row[at["Your call"]] or "").strip().lower()
            codes = {company: _code(row[at[column]]) for column, company in CODE_COLUMNS.items()}
            if call and call not in YES and call not in NO:
                problems.append(
                    f"#{transporter_id}: 'Your call' is {call!r}; use OK, manual or no "
                    "(a different code goes in the company's code column)"
                )
                continue
            if call in NO:
                skipped["marked manual/no"] += 1
                continue
            if match not in SURE_MATCHES and call not in YES:
                skipped[f"'{match}' not marked OK"] += 1
                continue
            if not any(codes.values()):
                skipped["no codes"] += 1
                continue
            plan.extend(
                (transporter_id, company, code) for company, code in codes.items() if code
            )
        return plan, skipped, problems
