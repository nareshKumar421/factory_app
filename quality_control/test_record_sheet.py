"""Tests for record forms uploaded as Excel sheets."""

import json
from datetime import date
from io import BytesIO

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Permission
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import SimpleTestCase
from django.urls import reverse
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from PIL import Image as PILImage
from rest_framework import status
from rest_framework.test import APIClient, APITestCase

from company.models import Company, UserCompany, UserRole
from quality_control.models import (
    QCRecord,
    RecordTemplate,
    RecordTemplateParameter,
    RecordTemplateSection,
)
from quality_control.services import record_sheet
from quality_control.services.record_sheet import FieldType

User = get_user_model()

THIN = Side(style="thin")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def build_form_workbook(with_image=True):
    """A small copy of the oil plant monitoring record's shape.

    Title over the top, 'Date:' at the right, a Sr No / Parameters / UOM /
    Time x3 table whose first body row is the time of each reading, then
    Remarks and the two signatures, with the controlled-document footer.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Product Testing "
    big = Font(name="Times New Roman", size=60, bold=True)

    ws["A1"] = "   OIL PLANT ON LINE MONITORING  RECORD"
    ws["A1"].font = big
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.merge_cells("A1:F1")
    ws["E2"] = "Date: "
    ws["E2"].font = big

    headings = {"A4": "Sr No.", "B4": "Parameters", "C4": "UOM", "D4": "Time", "E4": "Time", "F4": "Time"}
    for address, text in headings.items():
        ws[address] = text
    for row in (4, 5):
        for col in "ABCDEF":
            ws[f"{col}{row}"].border = BOX
            ws[f"{col}{row}"].font = big
    for block in ("A4:A5", "B4:B5", "C4:C5"):
        ws.merge_cells(block)
    ws["D4"].fill = PatternFill("solid", fgColor="FFD9D9D9")

    body = [
        (1, "PRODUCT", None),
        (2, "Free Fatty Acids ", "%"),
        (3, "Argemone Oil", "Absent"),
        (4, "TBHQ", "Present / Absent"),
    ]
    for offset, (number, name, unit) in enumerate(body):
        row = 6 + offset
        ws[f"A{row}"] = number
        ws[f"B{row}"] = name
        if unit:
            ws[f"C{row}"] = unit
        for col in "ABCDEF":
            ws[f"{col}{row}"].border = BOX
        ws.row_dimensions[row].height = 180

    ws["B11"] = "Remarks:"
    ws["F12"] = "Q.A.M"
    ws["B13"] = "Q.A Chemist"

    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 82.44
    ws.row_dimensions[1].height = 75

    ws.print_area = "A1:F14"
    ws.page_setup.orientation = "landscape"
    ws.oddHeader.center.text = "JIVO WELLNESS PVT.LTD."
    ws.oddFooter.left.text = "Revision No.:02/22-05-2026"
    ws.oddFooter.center.text = "Classified : Business Confidential "
    ws.oddFooter.right.text = "QA-FRM-14-01-05-02 Controlled Document"

    if with_image:
        buffer = BytesIO()
        PILImage.new("RGB", (40, 20), "red").save(buffer, "PNG")
        buffer.seek(0)
        ws.add_image(XLImage(buffer), "A2")

    out = BytesIO()
    wb.save(out)
    out.seek(0)
    return out


def upload(name="QA-FRM-14-01-05-02 ONLINE OIL PLANT MONITORING RECORD .xlsx", content=None):
    data = (content or build_form_workbook()).getvalue()
    return SimpleUploadedFile(
        name, data, content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )


# ---------------------------------------------------------------------------
# The parser
# ---------------------------------------------------------------------------


class ReadSheetTests(SimpleTestCase):
    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.result = record_sheet.read_sheet(build_form_workbook(), filename="form.xlsx")
        cls.layout = cls.result["layout"]
        cls.fields = cls.result["cell_fields"]

    def test_layout_follows_the_print_area(self):
        self.assertEqual(self.layout["range"], "A1:F14")
        self.assertEqual(len(self.layout["cols"]), 6)
        self.assertEqual(len(self.layout["rows"]), 14)
        self.assertEqual(self.layout["orientation"], "landscape")

    def test_sizes_are_converted_to_pixels(self):
        self.assertEqual(self.layout["cols"][0]["w"], 44 * 7 + 5)
        self.assertEqual(self.layout["rows"][0]["h"], 100)  # 75pt
        self.assertEqual(self.layout["rows"][5]["h"], 240)  # 180pt

    def test_cells_merges_and_styles(self):
        self.assertIn("A1:F1", self.layout["merges"])
        self.assertIn("A4:A5", self.layout["merges"])
        # Covered cells of a merge are not cells of their own.
        self.assertNotIn("B1", self.layout["cells"])
        title = self.layout["cells"]["A1"]
        self.assertEqual(title["v"], "   OIL PLANT ON LINE MONITORING  RECORD")
        style = self.layout["styles"][title["s"]]
        self.assertTrue(style["b"])
        self.assertEqual(style["fs"], 60)
        self.assertEqual(style["ha"], "center")
        self.assertEqual(style["va"], "middle")

        number = self.layout["cells"]["A6"]
        self.assertEqual(number["v"], "1")
        self.assertTrue(number["n"])
        boxed = self.layout["styles"][self.layout["cells"]["D6"]["s"]]
        self.assertEqual([boxed.get(k) for k in ("bl", "br", "bt", "bb")], ["thin"] * 4)
        filled = self.layout["styles"][self.layout["cells"]["D4"]["s"]]
        self.assertEqual(filled["bg"], "#d9d9d9")

    def test_merged_block_takes_its_edge_borders(self):
        # A4:A5 is boxed; the bottom edge comes from A5, not the anchor.
        style = self.layout["styles"][self.layout["cells"]["A4"]["s"]]
        self.assertEqual(style.get("bb"), "thin")

    def test_picture_and_header_footer(self):
        self.assertEqual(len(self.layout["images"]), 1)
        image = self.layout["images"][0]
        self.assertTrue(image["src"].startswith("data:image/png;base64,"))
        self.assertEqual(image["x"], 0)
        self.assertEqual(image["y"], 100)  # below the 75pt title row
        self.assertEqual(self.layout["header"]["center"], "JIVO WELLNESS PVT.LTD.")
        self.assertEqual(self.layout["footer"]["left"], "Revision No.:02/22-05-2026")

    def test_header_is_read_off_the_footer_and_sheet(self):
        self.assertEqual(
            self.result["header"],
            {
                "document_code": "QA-FRM-14-01-05-02",
                "title": "OIL PLANT ON LINE MONITORING RECORD",
                "organisation": "JIVO WELLNESS PVT.LTD.",
                "revision_number": "02",
                "revision_date": "2026-05-22",
                "classification": "Business Confidential",
            },
        )

    def test_readings_are_the_blank_boxes_of_blank_columns(self):
        readings = {cell for cell in self.fields if cell[0] in "DEF" and 5 <= int(cell[1:]) <= 9}
        self.assertEqual(len(readings), 15)
        # The UOM column has text in its boxes, so its blanks are not readings.
        self.assertNotIn("C6", self.fields)

    def test_reading_types_come_from_the_row(self):
        self.assertEqual(self.fields["D5"]["type"], FieldType.TIME)
        self.assertEqual(self.fields["F5"]["label"], "Time 3")
        self.assertEqual(self.fields["D6"]["type"], FieldType.TEXT)
        self.assertEqual(self.fields["D6"]["label"], "PRODUCT · Time 1")
        self.assertEqual(self.fields["E7"]["type"], FieldType.NUMBER)
        self.assertEqual(self.fields["D8"]["type"], FieldType.CHOICE)
        self.assertEqual(self.fields["D8"]["options"], ["Absent", "Present"])
        self.assertEqual(self.fields["D8"]["ok"], ["Absent"])
        # Either answer is acceptable, so nothing is judged.
        self.assertEqual(self.fields["D9"]["ok"], [])

    def test_labels_bind_their_neighbour_to_the_record(self):
        self.assertEqual(self.fields["F2"]["type"], FieldType.RECORD_DATE)
        self.assertEqual(self.fields["C11"]["type"], FieldType.REMARKS)
        self.assertEqual(self.fields["C13"]["type"], FieldType.SIGN_SUBMITTED)
        # Q.A.M sits in the last column, so its signature goes below it.
        self.assertEqual(self.fields["F13"]["type"], FieldType.SIGN_APPROVED)

    def test_other_sheets_can_be_chosen(self):
        wb = Workbook()
        wb.active.title = "Cover"
        wb.active["A1"] = "cover page"
        second = wb.create_sheet("Form")
        second["B2"] = "Batch No:"
        out = BytesIO()
        wb.save(out)
        out.seek(0)
        result = record_sheet.read_sheet(out, filename="x.xlsx", sheet_name="Form")
        self.assertEqual(result["sheets"], ["Cover", "Form"])
        self.assertEqual(result["sheet"], "Form")
        with self.assertRaises(record_sheet.SheetImportError):
            out.seek(0)
            record_sheet.read_sheet(out, filename="x.xlsx", sheet_name="Missing")


class UploadRefusalTests(SimpleTestCase):
    def test_old_xls_is_refused_with_a_way_forward(self):
        with self.assertRaisesMessage(record_sheet.SheetImportError, "save it as"):
            record_sheet.read_sheet(BytesIO(b"whatever"), filename="old.xls")

    def test_non_workbook_is_refused(self):
        with self.assertRaises(record_sheet.SheetImportError):
            record_sheet.read_sheet(BytesIO(b"not a zip"), filename="a.xlsx")

    def test_a_data_dump_is_not_a_form(self):
        wb = Workbook()
        wb.active.cell(row=500, column=2, value="far away")
        wb.active["A1"] = "top"
        out = BytesIO()
        wb.save(out)
        out.seek(0)
        with self.assertRaisesMessage(record_sheet.SheetImportError, "too big"):
            record_sheet.read_sheet(out, filename="a.xlsx")


class LayoutTokenTests(SimpleTestCase):
    def setUp(self):
        self.result = record_sheet.read_sheet(build_form_workbook(), filename="form.xlsx")

    def test_the_layout_survives_a_trip_through_the_browser(self):
        # JavaScript re-serialises 60.0 as 60; the token must still match.
        def as_javascript(value):
            if isinstance(value, float) and value.is_integer():
                return int(value)
            if isinstance(value, dict):
                return {k: as_javascript(v) for k, v in value.items()}
            if isinstance(value, list):
                return [as_javascript(v) for v in value]
            return value

        travelled = as_javascript(json.loads(json.dumps(self.result["layout"])))
        self.assertTrue(record_sheet.verify_layout(travelled, self.result["layout_token"]))

    def test_a_changed_layout_is_refused(self):
        layout = json.loads(json.dumps(self.result["layout"]))
        layout["cells"]["A1"]["v"] = "Something else"
        self.assertFalse(record_sheet.verify_layout(layout, self.result["layout_token"]))
        self.assertFalse(record_sheet.verify_layout(self.result["layout"], "forged"))
        self.assertFalse(record_sheet.verify_layout(self.result["layout"], ""))


class CellRulesTests(SimpleTestCase):
    def test_parse_specification(self):
        cases = {
            "6.5 - 8.5": ("6.5", "8.5"),
            "Max 2.0 NTU": (None, "2.0"),
            "NMT 0.5": (None, "0.5"),
            "NLT 20": ("20", None),
            "10 ± 2": ("8", "12"),
            "To be tested": (None, None),
            "": (None, None),
        }
        for text, expected in cases.items():
            self.assertEqual(record_sheet.parse_specification(text), expected, text)

    def test_check_cell(self):
        number = {"type": FieldType.NUMBER, "min": "6.5", "max": "8.5"}
        self.assertTrue(record_sheet.check_cell(number, "7"))
        self.assertFalse(record_sheet.check_cell(number, "9"))
        self.assertFalse(record_sheet.check_cell(number, "n/a"))
        self.assertIsNone(record_sheet.check_cell(number, ""))
        self.assertIsNone(record_sheet.check_cell({"type": FieldType.NUMBER}, "9"))

        choice = {"type": FieldType.CHOICE, "options": ["Absent", "Present"], "ok": ["Absent"]}
        self.assertTrue(record_sheet.check_cell(choice, "absent"))
        self.assertFalse(record_sheet.check_cell(choice, "Present"))
        self.assertIsNone(record_sheet.check_cell({"type": FieldType.CHOICE}, "Present"))
        self.assertIsNone(record_sheet.check_cell({"type": FieldType.TEXT}, "x"))

    def test_clean_cell_value(self):
        self.assertEqual(record_sheet.clean_cell_value({"type": FieldType.TIME}, "8:05"), "08:05")
        self.assertEqual(record_sheet.clean_cell_value({"type": FieldType.TIME}, "14:30:00"), "14:30")
        with self.assertRaises(ValueError):
            record_sheet.clean_cell_value({"type": FieldType.TIME}, "25:00")
        self.assertEqual(
            record_sheet.clean_cell_value({"type": FieldType.DATE}, "2026-09-26"), "2026-09-26"
        )
        with self.assertRaises(ValueError):
            record_sheet.clean_cell_value({"type": FieldType.DATE}, "26/09/2026")
        with self.assertRaises(ValueError):
            record_sheet.clean_cell_value({"type": FieldType.TEXT}, "x" * 256)
        self.assertEqual(record_sheet.clean_cell_value({"type": FieldType.TEXT}, "  "), "")

    def test_clean_cell_fields(self):
        layout = record_sheet.read_sheet(build_form_workbook(), filename="f.xlsx")["layout"]
        cleaned, errors = record_sheet.clean_cell_fields(
            {
                "D7": {"type": "NUMBER", "min": "0", "max": " 0.5 ", "junk": "dropped"},
                "D8": {"type": "CHOICE", "options": ["Absent", "Present", "Absent"], "ok": ["Absent"]},
            },
            layout,
        )
        self.assertEqual(errors, {})
        self.assertEqual(cleaned["D7"], {"type": "NUMBER", "min": "0", "max": "0.5"})
        self.assertEqual(cleaned["D8"]["options"], ["Absent", "Present"])

        _, errors = record_sheet.clean_cell_fields(
            {
                "B1": {"type": "TEXT"},  # inside the A1:F1 title
                "Z99": {"type": "TEXT"},  # off the sheet
                "D6": {"type": "SIGNATURE"},
                "D7": {"type": "NUMBER", "min": "5", "max": "1"},
                "C11": {"type": "REMARKS"},
                "C12": {"type": "REMARKS"},
            },
            layout,
        )
        self.assertIn("A1", errors["B1"])
        self.assertIn("outside", errors["Z99"])
        self.assertIn("type", errors["D6"])
        self.assertIn("above", errors["D7"])
        self.assertIn("remarks", errors["__all__"])


# ---------------------------------------------------------------------------
# The endpoints
# ---------------------------------------------------------------------------


def _client(company, perms="all"):
    n = User.objects.count()
    user = User.objects.create_user(
        email=f"s{n}@t.com", password="x", full_name=f"Sheet User {n}", employee_code=f"S{n}"
    )
    role = UserRole.objects.create(name=f"SR{UserRole.objects.count()}")
    UserCompany.objects.create(user=user, company=company, role=role, is_active=True)
    qs = Permission.objects.filter(content_type__app_label="quality_control")
    if perms != "all":
        qs = qs.filter(codename__in=perms)
    user.user_permissions.set(qs)
    client = APIClient()
    client.force_authenticate(user=User.objects.get(pk=user.pk))
    client.credentials(HTTP_COMPANY_CODE=company.code)
    return client


class SheetApiTestBase(APITestCase):
    def setUp(self):
        self.company = Company.objects.create(code="SHEET_CO", name="Sheet Co")
        self.client = _client(self.company)
        response = self.client.post(
            reverse("record-template-import-sheet"), {"file": upload()}, format="multipart"
        )
        self.assertEqual(response.status_code, status.HTTP_200_OK, response.data)
        self.preview = response.data

    def create_template(self, **overrides):
        payload = {
            **self.preview["header"],
            "layout": self.preview["layout"],
            "layout_token": self.preview["layout_token"],
            "cell_fields": self.preview["cell_fields"],
            "source_file_name": self.preview["source_file_name"],
            **overrides,
        }
        return self.client.post(reverse("record-template-list-create"), payload, format="json")

    def open_record(self, template_id, day=date(2026, 9, 26)):
        response = self.client.post(
            reverse("qc-record-list-create"),
            {"template": template_id, "record_date": day.isoformat()},
            format="json",
        )
        self.assertIn(response.status_code, (200, 201), response.data)
        return response.data["id"]

    def save_cells(self, record_id, cells, **extra):
        return self.client.post(
            reverse("qc-record-cells", args=[record_id]), {"cells": cells, **extra}, format="json"
        )


class SheetImportApiTests(SheetApiTestBase):
    def test_preview_carries_everything_the_designer_needs(self):
        self.assertEqual(self.preview["header"]["document_code"], "QA-FRM-14-01-05-02")
        self.assertEqual(self.preview["sheet"], "Product Testing ")
        self.assertTrue(self.preview["layout_token"])
        self.assertEqual(self.preview["cell_fields"]["D8"]["type"], "CHOICE")
        self.assertTrue(self.preview["source_file_name"].endswith(".xlsx"))
        # Nothing is saved by a preview.
        self.assertFalse(RecordTemplate.objects.exists())

    def test_only_a_form_maintainer_may_upload(self):
        filler = _client(self.company, perms=["can_view_qc_records", "can_fill_qc_records"])
        response = filler.post(
            reverse("record-template-import-sheet"), {"file": upload()}, format="multipart"
        )
        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)

    def test_unreadable_upload_is_a_400_with_a_reason(self):
        response = self.client.post(
            reverse("record-template-import-sheet"),
            {"file": SimpleUploadedFile("old.xls", b"x")},
            format="multipart",
        )
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn(".xlsx", response.data["detail"])
        response = self.client.post(reverse("record-template-import-sheet"), {}, format="multipart")
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)


class SheetTemplateApiTests(SheetApiTestBase):
    def test_create_a_sheet_form(self):
        response = self.create_template(document_code="qa-frm-14-01-05-02")
        self.assertEqual(response.status_code, status.HTTP_201_CREATED, response.data)
        self.assertEqual(response.data["kind"], "SHEET")
        self.assertEqual(response.data["document_code"], "QA-FRM-14-01-05-02")
        self.assertNotIn("layout_token", response.data)
        template = RecordTemplate.objects.get(pk=response.data["id"])
        self.assertIsNone(template.company)  # shared, like every form
        self.assertEqual(template.layout["range"], "A1:F14")
        self.assertEqual(template.cell_fields["F2"]["type"], "RECORD_DATE")
        # Empty option lists from the guess are not stored.
        self.assertNotIn("ok", template.cell_fields["D9"])

        listing = self.client.get(reverse("record-template-list-create")).data
        self.assertEqual(listing[0]["kind"], "SHEET")
        self.assertEqual(listing[0]["field_count"], 15)
        self.assertNotIn("layout", listing[0])

    def test_a_layout_needs_the_token_the_import_issued(self):
        layout = json.loads(json.dumps(self.preview["layout"]))
        layout["cells"]["A1"]["v"] = "Forged"
        response = self.create_template(layout=layout)
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("layout", response.data)
        response = self.create_template(layout_token="")
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)

    def test_bad_cell_fields_are_refused_per_cell(self):
        response = self.create_template(cell_fields={"B1": {"type": "TEXT"}})
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("B1", response.data["cell_fields"])

    def test_a_parameter_form_cannot_become_a_sheet(self):
        grid = RecordTemplate.objects.create(document_code="GRID-1", title="Grid form")
        response = self.client.put(
            reverse("record-template-detail", args=[grid.id]),
            {"layout": self.preview["layout"], "layout_token": self.preview["layout_token"]},
            format="json",
        )
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)

    def test_filled_form_keeps_its_cells_but_not_its_header(self):
        template_id = self.create_template().data["id"]
        record_id = self.open_record(template_id)
        self.assertEqual(self.save_cells(record_id, {"D7": "0.12"}).status_code, 200)
        url = reverse("record-template-detail", args=[template_id])

        detail = self.client.get(url).data
        self.assertTrue(detail["is_locked"])

        changed = dict(detail["cell_fields"])
        changed["D7"] = {"type": "NUMBER", "max": "0.1"}
        response = self.client.put(url, {"cell_fields": changed}, format="json")
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)

        # Re-sending the form's own fields with a header fix is fine.
        response = self.client.put(
            url,
            {"revision_number": "03", "cell_fields": detail["cell_fields"]},
            format="json",
        )
        self.assertEqual(response.status_code, status.HTTP_200_OK, response.data)
        self.assertEqual(response.data["revision_number"], "03")

    def test_an_unfilled_form_can_be_reworked(self):
        template_id = self.create_template().data["id"]
        self.open_record(template_id)  # opened, but nothing typed yet
        fields = dict(self.preview["cell_fields"])
        fields["D7"] = {"type": "NUMBER", "min": "0", "max": "0.5"}
        response = self.client.put(
            reverse("record-template-detail", args=[template_id]),
            {"cell_fields": fields},
            format="json",
        )
        self.assertEqual(response.status_code, status.HTTP_200_OK, response.data)
        self.assertEqual(response.data["cell_fields"]["D7"]["max"], "0.5")


class SheetRecordApiTests(SheetApiTestBase):
    def setUp(self):
        super().setUp()
        fields = dict(self.preview["cell_fields"])
        fields["E7"] = {"type": "NUMBER", "min": "0", "max": "0.5", "label": "FFA"}
        response = self.create_template(cell_fields=fields)
        self.assertEqual(response.status_code, status.HTTP_201_CREATED, response.data)
        self.template_id = response.data["id"]
        self.record_id = self.open_record(self.template_id)

    def test_cells_are_saved_merged_and_judged(self):
        response = self.save_cells(
            self.record_id,
            {"D5": "8:10", "D6": "Olive Oil", "E7": "0.9", "D8": "Present"},
            remarks="Line 2 FFA high",
        )
        self.assertEqual(response.status_code, status.HTTP_200_OK, response.data)
        data = response.data
        self.assertEqual(data["cell_values"]["D5"], "08:10")
        self.assertEqual(data["remarks"], "Line 2 FFA high")
        self.assertEqual(data["cell_checks"], {"E7": False, "D8": False})

        # A second save touches only what it sends; blank clears a cell.
        response = self.save_cells(self.record_id, {"E7": "0.2", "D6": ""})
        values = response.data["cell_values"]
        self.assertEqual(values, {"D5": "08:10", "E7": "0.2", "D8": "Present"})
        self.assertEqual(response.data["cell_checks"]["E7"], True)
        self.assertEqual(response.data["remarks"], "Line 2 FFA high")

        listing = self.client.get(reverse("qc-record-list-create")).data
        row = next(item for item in listing if item["id"] == self.record_id)
        self.assertEqual(row["filled_count"], 3)
        self.assertEqual(row["slot_count"], 1)

    def test_only_fillable_cells_are_accepted(self):
        response = self.save_cells(self.record_id, {"A1": "x", "F2": "2026-09-26"})
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(set(response.data["cells"]), {"A1", "F2"})
        response = self.save_cells(self.record_id, {"D5": "late"})
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("time", response.data["cells"]["D5"])
        self.assertEqual(QCRecord.objects.get(pk=self.record_id).cell_values, {})

    def test_approved_sheet_is_locked(self):
        self.save_cells(self.record_id, {"D6": "Olive Oil"})
        self.client.post(reverse("qc-record-submit", args=[self.record_id]))
        response = self.client.post(
            reverse("qc-record-approve", args=[self.record_id]), {"decision": "APPROVE"}, format="json"
        )
        self.assertEqual(response.data["status"], "APPROVED")
        response = self.save_cells(self.record_id, {"D6": "Canola"})
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)

    def test_the_two_kinds_of_form_keep_to_their_own_save(self):
        response = self.client.post(
            reverse("qc-record-values", args=[self.record_id]),
            {"cells": [{"slot_time": "08:00", "parameter": 1, "value": "x"}]},
            format="json",
        )
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)

        grid = RecordTemplate.objects.create(document_code="GRID-2", title="Grid form")
        section = RecordTemplateSection.objects.create(template=grid, title="S")
        RecordTemplateParameter.objects.create(section=section, name="pH")
        grid_record = self.open_record(grid.id)
        response = self.save_cells(grid_record, {"D6": "x"})
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)

    def test_detail_carries_the_sheet(self):
        data = self.client.get(reverse("qc-record-detail", args=[self.record_id])).data
        self.assertEqual(data["template_detail"]["kind"], "SHEET")
        self.assertEqual(data["template_detail"]["layout"]["range"], "A1:F14")
        self.assertEqual(data["cell_values"], {})
