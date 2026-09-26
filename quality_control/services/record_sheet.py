"""
Record forms uploaded as Excel sheets.

QA keeps its record forms as Excel workbooks, laid out to print on one page.
Rather than re-typing each one into the parameter builder, the manager uploads
the workbook and the form is drawn on screen exactly as the sheet looks: the
same merges, borders, fonts, column widths and logo.

Three things come out of an upload:

* a **layout** -- the sheet itself, as data. It is only ever produced here, from
  the workbook, and is handed back to the browser with a signed token so a save
  can prove the layout it carries is the one this parser made
  (:func:`sign_layout` / :func:`verify_layout`).
* **cell fields** -- which cells are filled in, and how: a number, a time, one
  of "Absent / Present", the record's date, a signature. These are a first
  guess (:func:`detect_fields`) that the manager corrects before saving.
* a **header** -- the document code, revision, classification and title, read
  off the print footer and the sheet where they can be found.

A filled sheet is then just ``{cell: value}`` against that layout.
"""

import base64
import colorsys
import hashlib
import json
import re
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation
from xml.etree import ElementTree

from django.core import signing
from openpyxl import load_workbook
from openpyxl.styles.colors import COLOR_INDEX
from openpyxl.utils.cell import (
    column_index_from_string,
    coordinate_from_string,
    get_column_letter,
    range_boundaries,
)
from openpyxl.utils.exceptions import InvalidFileException

LAYOUT_VERSION = 1

MAX_UPLOAD_BYTES = 5 * 1024 * 1024
# A record form prints on a page or two. Anything bigger is a data sheet, not
# a form, and would only produce an unusable screen.
MAX_COLUMNS = 80
MAX_ROWS = 400
MAX_CELLS = 12000
MAX_IMAGE_BYTES = 1024 * 1024
MAX_IMAGES_BYTES = 3 * 1024 * 1024

# How long an uploaded layout may sit in the designer before it must be saved.
LAYOUT_TOKEN_MAX_AGE = 24 * 60 * 60
_LAYOUT_SALT = "quality_control.record_sheet.layout"


class FieldType:
    """What a filled-in cell holds.

    The first five are values typed into the sheet and stored per record. The
    rest are *bound* to the record itself, so they are shown, never typed:
    the date the sheet is for, its shift, its remarks, and who submitted and
    approved it (the Q.A Chemist and Q.A.M signatures on the paper form).
    """

    TEXT = "TEXT"
    NUMBER = "NUMBER"
    TIME = "TIME"
    DATE = "DATE"
    CHOICE = "CHOICE"

    RECORD_DATE = "RECORD_DATE"
    SHIFT = "SHIFT"
    REMARKS = "REMARKS"
    SIGN_SUBMITTED = "SIGN_SUBMITTED"
    SIGN_APPROVED = "SIGN_APPROVED"


VALUE_TYPES = {
    FieldType.TEXT,
    FieldType.NUMBER,
    FieldType.TIME,
    FieldType.DATE,
    FieldType.CHOICE,
}
BOUND_TYPES = {
    FieldType.RECORD_DATE,
    FieldType.SHIFT,
    FieldType.REMARKS,
    FieldType.SIGN_SUBMITTED,
    FieldType.SIGN_APPROVED,
}
FIELD_TYPES = VALUE_TYPES | BOUND_TYPES

MAX_VALUE_LENGTH = 255


class SheetImportError(Exception):
    """The upload cannot be turned into a form; the message says why."""


# ---------------------------------------------------------------------------
# Cell references
# ---------------------------------------------------------------------------


def ref(row, col):
    return f"{get_column_letter(col)}{row}"


def parse_ref(value):
    """'D6' -> (6, 4). Raises ValueError for anything that is not one cell."""
    if not isinstance(value, str) or not re.fullmatch(r"[A-Z]{1,3}[1-9][0-9]{0,6}", value):
        raise ValueError(f"'{value}' is not a cell reference.")
    letters, row = coordinate_from_string(value)
    return row, column_index_from_string(letters)


# ---------------------------------------------------------------------------
# Colours
# ---------------------------------------------------------------------------

# Office's default palette, for a workbook that carries no theme of its own.
_DEFAULT_THEME = [
    "FFFFFF", "000000", "E7E6E6", "44546A", "4472C4", "ED7D31",
    "A5A5A5", "FFC000", "5B9BD5", "70AD47", "0563C1", "954F72",
]


def _theme_palette(workbook):
    """The workbook's theme colours, in the order styles index them.

    The theme XML lists dk1, lt1, dk2, lt2, accents...; a style's
    ``theme=0`` means lt1 and ``theme=1`` dk1 -- the first two pairs swap.
    """
    raw = getattr(workbook, "loaded_theme", None)
    if not raw:
        return list(_DEFAULT_THEME)
    try:
        root = ElementTree.fromstring(raw)
    except ElementTree.ParseError:
        return list(_DEFAULT_THEME)
    ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    scheme = root.find(".//a:clrScheme", ns)
    if scheme is None:
        return list(_DEFAULT_THEME)
    names = [
        "lt1", "dk1", "lt2", "dk2", "accent1", "accent2", "accent3",
        "accent4", "accent5", "accent6", "hlink", "folHlink",
    ]
    palette = []
    for index, name in enumerate(names):
        element = scheme.find(f"a:{name}", ns)
        value = None
        if element is not None and len(element):
            child = element[0]
            value = child.get("val") if child.tag.endswith("srgbClr") else child.get("lastClr")
        palette.append((value or _DEFAULT_THEME[index]).upper())
    return palette


def _tint(hex6, tint):
    """Apply Excel's tint (-1 darker ... +1 lighter) to an RRGGBB colour."""
    if not tint:
        return hex6
    red, green, blue = (int(hex6[i : i + 2], 16) / 255 for i in (0, 2, 4))
    hue, light, sat = colorsys.rgb_to_hls(red, green, blue)
    light = light * (1 + tint) if tint < 0 else light * (1 - tint) + tint
    red, green, blue = colorsys.hls_to_rgb(hue, max(0.0, min(1.0, light)), sat)
    return "".join(f"{round(channel * 255):02X}" for channel in (red, green, blue))


def _colour(colour, palette):
    """An openpyxl colour as '#rrggbb', or None when it is unset."""
    if colour is None:
        return None
    kind = getattr(colour, "type", None)
    hex6 = None
    if kind == "rgb":
        value = colour.rgb
        if isinstance(value, str) and re.fullmatch(r"[0-9A-Fa-f]{8}", value):
            hex6 = value[2:]
    elif kind == "theme":
        index = colour.theme
        if isinstance(index, int) and 0 <= index < len(palette):
            hex6 = _tint(palette[index], colour.tint or 0)
    elif kind == "indexed":
        index = colour.indexed
        # 64/65 are "system foreground/background": the default, not a colour.
        if isinstance(index, int) and 0 <= index < len(COLOR_INDEX) and index < 64:
            hex6 = COLOR_INDEX[index][2:]
    return f"#{hex6.lower()}" if hex6 else None


# ---------------------------------------------------------------------------
# Sizes
# ---------------------------------------------------------------------------


def _column_px(width):
    """Excel column width (characters of Calibri 11) to screen pixels."""
    return round(width * 7 + 5, 2)


def _points_px(points):
    return round(points * 96 / 72, 2)


# ---------------------------------------------------------------------------
# Cell values and styles
# ---------------------------------------------------------------------------


def _display(cell):
    """The text Excel would show in a cell, as far as a form ever needs."""
    value = cell.value
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, (int, float)):
        fmt = cell.number_format or "General"
        decimals = re.search(r"0\.(0+)", fmt)
        if "%" in fmt:
            places = len(decimals.group(1)) if decimals else 0
            return f"{value * 100:.{places}f}%"
        if decimals:
            return f"{value:.{len(decimals.group(1))}f}"
        if isinstance(value, float):
            return str(int(value)) if value.is_integer() else f"{value:.10g}"
        return str(value)
    if isinstance(value, datetime):
        if value.time() == time(0, 0):
            return value.strftime("%d-%m-%Y")
        return value.strftime("%d-%m-%Y %H:%M")
    if isinstance(value, date):
        return value.strftime("%d-%m-%Y")
    if isinstance(value, time):
        return value.strftime("%H:%M")
    return str(value).replace("\r\n", "\n").replace("\r", "\n")


_BORDER_STYLES = {
    "thin", "medium", "thick", "dashed", "dotted", "double", "hair",
    "mediumDashed", "dashDot", "mediumDashDot", "dashDotDot",
    "mediumDashDotDot", "slantDashDot",
}
_H_ALIGN = {"left", "center", "right", "justify", "centerContinuous", "distributed", "fill"}
_V_ALIGN = {"top", "center", "bottom", "justify", "distributed"}


def _border_side(side):
    style = getattr(side, "style", None)
    return style if style in _BORDER_STYLES else None


def _style(cell, palette, borders=None):
    """The parts of a cell's look the screen reproduces, in compact keys."""
    font = cell.font
    alignment = cell.alignment
    fill = cell.fill
    style = {}

    if font is not None:
        if font.name:
            style["ff"] = str(font.name)[:64]
        if font.sz:
            style["fs"] = float(font.sz)
        if font.b:
            style["b"] = True
        if font.i:
            style["i"] = True
        if font.u and font.u != "none":
            style["u"] = True
        if font.strike:
            style["st"] = True
        colour = _colour(font.color, palette)
        if colour and colour != "#000000":
            style["fc"] = colour

    if fill is not None and fill.fill_type and fill.fill_type != "none":
        colour = _colour(fill.fgColor, palette) or _colour(fill.start_color, palette)
        if colour:
            style["bg"] = colour

    if alignment is not None:
        if alignment.horizontal in _H_ALIGN:
            horizontal = alignment.horizontal
            style["ha"] = {"centerContinuous": "center", "distributed": "justify", "fill": "left"}.get(
                horizontal, horizontal
            )
        if alignment.vertical in _V_ALIGN:
            style["va"] = {"center": "middle", "justify": "middle", "distributed": "middle"}.get(
                alignment.vertical, alignment.vertical
            )
        if alignment.wrap_text:
            style["wr"] = True
        if alignment.text_rotation:
            style["rot"] = int(alignment.text_rotation)
        if alignment.indent:
            style["ind"] = int(alignment.indent)

    border = borders if borders is not None else cell.border
    for key, side in (("bl", "left"), ("br", "right"), ("bt", "top"), ("bb", "bottom")):
        value = _border_side(getattr(border, side, None)) if border is not None else None
        if value:
            style[key] = value
    return style


class _MergedBorder:
    """The outer edges of a merged block, taken from the cells on each edge."""

    def __init__(self, worksheet, r1, c1, r2, c2):
        self.left = worksheet.cell(r1, c1).border.left
        self.top = worksheet.cell(r1, c1).border.top
        self.right = worksheet.cell(r1, c2).border.right
        self.bottom = worksheet.cell(r2, c1).border.bottom


# ---------------------------------------------------------------------------
# Range, images, header/footer
# ---------------------------------------------------------------------------


def _print_range(worksheet):
    """The sheet's print area if it has one, else the part actually used."""
    area = worksheet.print_area
    if area:
        first = str(area).split(",")[0].split("!")[-1].replace("$", "").strip("'")
        try:
            min_col, min_row, max_col, max_row = range_boundaries(first)
            if min_col and min_row and max_col and max_row:
                return min_row, min_col, max_row, max_col
        except (TypeError, ValueError):
            pass

    bounds = None

    def grow(row, col):
        nonlocal bounds
        if bounds is None:
            bounds = [row, col, row, col]
        else:
            bounds = [min(bounds[0], row), min(bounds[1], col), max(bounds[2], row), max(bounds[3], col)]

    for row in worksheet.iter_rows():
        for cell in row:
            if not hasattr(cell, "column"):
                continue
            has_value = cell.value is not None and str(cell.value).strip() != ""
            border = cell.border
            has_border = border is not None and any(
                _border_side(getattr(border, side, None)) for side in ("left", "right", "top", "bottom")
            )
            has_fill = cell.fill is not None and cell.fill.fill_type not in (None, "none")
            if has_value or has_border or has_fill:
                grow(cell.row, cell.column)
    for merged in worksheet.merged_cells.ranges:
        grow(merged.min_row, merged.min_col)
        grow(merged.max_row, merged.max_col)
    if bounds is None:
        raise SheetImportError("This sheet is empty.")
    return tuple(bounds)


def _anchor_position(marker, col_starts, row_starts, min_row, min_col):
    """Pixel position of a drawing anchor marker, relative to the range."""
    col = marker.col + 1  # markers are zero-based
    row = marker.row + 1
    if col < min_col or row < min_row:
        return None
    col_index = col - min_col
    row_index = row - min_row
    if col_index >= len(col_starts) or row_index >= len(row_starts):
        return None
    x = col_starts[col_index] + (marker.colOff or 0) / 9525
    y = row_starts[row_index] + (marker.rowOff or 0) / 9525
    return x, y


def _images(worksheet, col_starts, row_starts, min_row, min_col):
    """Pictures (a logo, a stamp) as data URLs placed over the grid."""
    placed = []
    total = 0
    for image in getattr(worksheet, "_images", []):
        anchor = image.anchor
        start = getattr(anchor, "_from", None)
        if start is None:
            continue
        origin = _anchor_position(start, col_starts, row_starts, min_row, min_col)
        if origin is None:
            continue
        end = getattr(anchor, "to", None)
        extent = getattr(anchor, "ext", None)
        if end is not None:
            corner = _anchor_position(end, col_starts, row_starts, min_row, min_col)
            if corner is None:
                continue
            width, height = corner[0] - origin[0], corner[1] - origin[1]
        elif extent is not None:
            width = (extent.width if hasattr(extent, "width") else extent.cx) / 9525
            height = (extent.height if hasattr(extent, "height") else extent.cy) / 9525
        else:
            width, height = image.width, image.height
        if width <= 0 or height <= 0:
            continue

        try:
            data = image._data()
        except Exception:  # an unreadable picture is not worth failing the upload
            continue
        if len(data) > MAX_IMAGE_BYTES or total + len(data) > MAX_IMAGES_BYTES:
            continue
        total += len(data)
        kind = (getattr(image, "format", None) or "png").lower()
        kind = {"jpg": "jpeg"}.get(kind, kind)
        if kind not in {"png", "jpeg", "gif", "bmp", "webp"}:
            continue
        placed.append(
            {
                "src": f"data:image/{kind};base64,{base64.b64encode(data).decode('ascii')}",
                "x": round(origin[0], 2),
                "y": round(origin[1], 2),
                "w": round(width, 2),
                "h": round(height, 2),
            }
        )
    return placed


def _header_footer(item):
    parts = {}
    for position in ("left", "center", "right"):
        part = getattr(item, position, None)
        text = getattr(part, "text", None) if part is not None else None
        parts[position] = (text or "").strip()
    return parts


# ---------------------------------------------------------------------------
# Layout
# ---------------------------------------------------------------------------


def build_layout(worksheet, palette):
    """The sheet as data: sizes, cells (value + style), merges, pictures."""
    min_row, min_col, max_row, max_col = _print_range(worksheet)
    n_cols = max_col - min_col + 1
    n_rows = max_row - min_row + 1
    if n_cols > MAX_COLUMNS or n_rows > MAX_ROWS or n_cols * n_rows > MAX_CELLS:
        raise SheetImportError(
            f"This sheet spans {n_cols} columns x {n_rows} rows, which is too big for "
            "a record form. Set a print area around the form in Excel and upload it again."
        )

    default_width = worksheet.sheet_format.defaultColWidth or (
        (worksheet.sheet_format.baseColWidth or 8) + 0.43
    )
    widths = {}
    hidden_cols = set()
    for dimension in worksheet.column_dimensions.values():
        low = dimension.min or column_index_from_string(dimension.index)
        high = dimension.max or low
        for col in range(max(low, min_col), min(high, max_col) + 1):
            if dimension.width:
                widths[col] = dimension.width
            if dimension.hidden:
                hidden_cols.add(col)

    cols = []
    for col in range(min_col, max_col + 1):
        entry = {"w": _column_px(widths.get(col, default_width))}
        if col in hidden_cols:
            entry["hidden"] = True
        cols.append(entry)

    merges = []
    merge_of = {}  # anchor ref -> (r1, c1, r2, c2)
    covered = set()
    for merged in worksheet.merged_cells.ranges:
        r1, c1 = max(merged.min_row, min_row), max(merged.min_col, min_col)
        r2, c2 = min(merged.max_row, max_row), min(merged.max_col, max_col)
        if r1 > r2 or c1 > c2 or (r1 == r2 and c1 == c2):
            continue
        # Only a block whose anchor survives clipping keeps its content.
        if (r1, c1) != (merged.min_row, merged.min_col):
            continue
        merges.append(f"{ref(r1, c1)}:{ref(r2, c2)}")
        merge_of[ref(r1, c1)] = (r1, c1, r2, c2)
        for row in range(r1, r2 + 1):
            for col in range(c1, c2 + 1):
                if (row, col) != (r1, c1):
                    covered.add(ref(row, col))

    styles = []
    style_index = {}
    cells = {}
    tallest_font = {}
    for row in range(min_row, max_row + 1):
        for col in range(min_col, max_col + 1):
            address = ref(row, col)
            if address in covered:
                continue
            cell = worksheet.cell(row, col)
            block = merge_of.get(address)
            borders = _MergedBorder(worksheet, *block) if block else None
            style = _style(cell, palette, borders)
            key = json.dumps(style, sort_keys=True)
            if key not in style_index:
                style_index[key] = len(styles)
                styles.append(style)
            entry = {"s": style_index[key]}
            text = _display(cell)
            if text != "":
                entry["v"] = text
                if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                    entry["n"] = True
                if not block:
                    tallest_font[row] = max(tallest_font.get(row, 0), style.get("fs", 11))
            cells[address] = entry

    default_height = worksheet.sheet_format.defaultRowHeight or 15
    rows = []
    for row in range(min_row, max_row + 1):
        dimension = worksheet.row_dimensions.get(row)
        height = dimension.ht if dimension is not None and dimension.ht else None
        if height is None:
            # Excel grows an unsized row to fit its biggest font.
            height = max(default_height, tallest_font.get(row, 0) * 1.3)
        entry = {"h": _points_px(height)}
        if dimension is not None and dimension.hidden:
            entry["hidden"] = True
        rows.append(entry)

    col_starts, running = [], 0.0
    for entry in cols:
        col_starts.append(running)
        running += 0 if entry.get("hidden") else entry["w"]
    row_starts, running = [], 0.0
    for entry in rows:
        row_starts.append(running)
        running += 0 if entry.get("hidden") else entry["h"]

    orientation = getattr(worksheet.page_setup, "orientation", None)
    return normalise({
        "version": LAYOUT_VERSION,
        "sheet": worksheet.title,
        "range": f"{ref(min_row, min_col)}:{ref(max_row, max_col)}",
        "cols": cols,
        "rows": rows,
        "styles": styles,
        "cells": cells,
        "merges": merges,
        "images": _images(worksheet, col_starts, row_starts, min_row, min_col),
        "header": _header_footer(worksheet.oddHeader),
        "footer": _header_footer(worksheet.oddFooter),
        "orientation": orientation if orientation in ("portrait", "landscape") else "portrait",
    })


# ---------------------------------------------------------------------------
# Signing: the layout a save carries must be one this parser produced
# ---------------------------------------------------------------------------


def normalise(value):
    """Whole-number floats as ints, recursively.

    A browser re-serialises ``60.0`` as ``60``, so a layout that went through
    the designer would otherwise no longer hash to its own token.
    """
    if isinstance(value, float) and value.is_integer():
        return int(value)
    if isinstance(value, dict):
        return {key: normalise(item) for key, item in value.items()}
    if isinstance(value, list):
        return [normalise(item) for item in value]
    return value


def _layout_digest(layout):
    canonical = json.dumps(
        normalise(layout), sort_keys=True, separators=(",", ":"), ensure_ascii=False
    )
    return hashlib.sha256(canonical.encode("utf-8")).hexdigest()


def sign_layout(layout):
    return signing.dumps(_layout_digest(layout), salt=_LAYOUT_SALT, compress=False)


def verify_layout(layout, token):
    """True when ``token`` was issued by :func:`sign_layout` for this layout."""
    if not isinstance(layout, dict) or not isinstance(token, str) or not token:
        return False
    try:
        digest = signing.loads(token, salt=_LAYOUT_SALT, max_age=LAYOUT_TOKEN_MAX_AGE)
    except signing.BadSignature:  # also covers SignatureExpired
        return False
    return digest == _layout_digest(layout)


# ---------------------------------------------------------------------------
# Field detection
# ---------------------------------------------------------------------------

_TIME_LABEL = re.compile(r"^\s*time\s*[:.\-]?\s*$", re.I)
_DATE_LABEL = re.compile(r"^\s*date\s*[:.\-]?\s*$", re.I)
_SHIFT_LABEL = re.compile(r"^\s*shift\s*[:.\-]?\s*$", re.I)
_REMARKS_LABEL = re.compile(r"^\s*remarks?\s*[:.\-]?\s*$", re.I)
_SUBMITTER_LABEL = re.compile(
    r"chemist|prepared\s*by|recorded\s*by|checked\s*by|done\s*by|tested\s*by|analy[sz]ed\s*by|filled\s*by",
    re.I,
)
_APPROVER_LABEL = re.compile(
    r"\bq\.?\s*a\.?\s*m\b\.?|manager|approved\s*by|verified\s*by|reviewed\s*by|authori[sz]ed\s*by|\bh\.?\s*o\.?\s*d\b",
    re.I,
)

_NAME_HEADER = re.compile(r"param|test|characteristic|description|particular|check\s*point|attribute|item", re.I)
_UNIT_HEADER = re.compile(r"^\s*(u\.?\s*o\.?\s*m\.?|units?|unit of measure(ment)?)\s*$", re.I)
_SPEC_HEADER = re.compile(r"spec|standard|limit|norm|acceptance|target|range", re.I)

_NUMBER = r"[-+]?\d+(?:\.\d+)?"


def _choice_from(text):
    """Options for an observation written as e.g. 'Absent' or 'Present / Absent'."""
    value = re.sub(r"\s+", " ", (text or "").strip().lower())
    value = re.sub(r"\s*/\s*", "/", value)
    pairs = {
        "absent": (["Absent", "Present"], ["Absent"]),
        "nil": (["Nil", "Present"], ["Nil"]),
        "present": (["Present", "Absent"], ["Present"]),
        "present/absent": (["Present", "Absent"], []),
        "absent/present": (["Absent", "Present"], []),
        "ok": (["OK", "Not OK"], ["OK"]),
        "ok/not ok": (["OK", "Not OK"], []),
        "yes/no": (["Yes", "No"], []),
        "pass/fail": (["Pass", "Fail"], []),
        "acceptable": (["Acceptable", "Not Acceptable"], ["Acceptable"]),
        "agreeable": (["Agreeable", "Not Agreeable"], ["Agreeable"]),
        "acceptable/not acceptable": (["Acceptable", "Not Acceptable"], []),
    }
    return pairs.get(value)


def parse_specification(text):
    """'6.5 - 8.5' -> ('6.5', '8.5'); 'Max 2.0' -> (None, '2.0'); else (None, None)."""
    value = (text or "").strip()
    if not value:
        return None, None
    match = re.search(rf"({_NUMBER})\s*(?:±|\+/-|\+-)\s*({_NUMBER})", value)
    if match:
        centre, spread = Decimal(match.group(1)), Decimal(match.group(2))
        return str(centre - spread), str(centre + spread)
    match = re.search(rf"^\s*({_NUMBER})\s*(?:-|–|—|to)\s*({_NUMBER})", value, re.I)
    if match:
        return match.group(1), match.group(2)
    match = re.search(
        rf"(?:max(?:imum)?\.?|nmt|not more than|≤|<=|<|up\s*to|below)\s*({_NUMBER})", value, re.I
    )
    if match:
        return None, match.group(1)
    match = re.search(
        rf"(?:min(?:imum)?\.?|nlt|not less than|≥|>=|>|above)\s*({_NUMBER})", value, re.I
    )
    if match:
        return match.group(1), None
    return None, None


class _Grid:
    """Read-only helpers over a layout for the detector."""

    def __init__(self, layout):
        self.layout = layout
        self.cells = layout["cells"]
        self.styles = layout["styles"]
        rng = layout["range"].split(":")
        self.min_row, self.min_col = parse_ref(rng[0])
        self.max_row, self.max_col = parse_ref(rng[1])
        self.merge_of = {}
        self.anchor_of = {}
        for block in layout["merges"]:
            start, end = block.split(":")
            r1, c1 = parse_ref(start)
            r2, c2 = parse_ref(end)
            self.merge_of[start] = (r1, c1, r2, c2)
            for row in range(r1, r2 + 1):
                for col in range(c1, c2 + 1):
                    if (row, col) != (r1, c1):
                        self.anchor_of[ref(row, col)] = start

    def inside(self, row, col):
        return self.min_row <= row <= self.max_row and self.min_col <= col <= self.max_col

    def covered(self, row, col):
        return ref(row, col) in self.anchor_of

    def text(self, row, col, follow_merge=True):
        address = ref(row, col)
        if address in self.anchor_of:
            if not follow_merge:
                return ""
            address = self.anchor_of[address]
        return (self.cells.get(address, {}).get("v") or "").strip()

    def style(self, row, col):
        entry = self.cells.get(ref(row, col))
        return self.styles[entry["s"]] if entry else {}

    def span(self, row, col):
        """(r2, c2) of the block a cell anchors -- itself when unmerged."""
        block = self.merge_of.get(ref(row, col))
        return (block[2], block[3]) if block else (row, col)

    def edge(self, row, col, side):
        """Whether a cell edge is drawn, by the cell itself or its neighbour."""
        own = {"left": "bl", "right": "br", "top": "bt", "bottom": "bb"}[side]
        if self.style(row, col).get(own):
            return True
        r2, c2 = self.span(row, col)
        neighbour = {
            "left": (row, col - 1, "br"),
            "right": (row, c2 + 1, "bl"),
            "top": (row - 1, col, "bb"),
            "bottom": (r2 + 1, col, "bt"),
        }[side]
        n_row, n_col, key = neighbour
        if not self.inside(n_row, n_col):
            return False
        address = ref(n_row, n_col)
        if address in self.anchor_of:
            anchor_row, anchor_col = parse_ref(self.anchor_of[address])
            return bool(self.style(anchor_row, anchor_col).get(key))
        return bool(self.style(n_row, n_col).get(key))

    def edges(self, row, col):
        return sum(self.edge(row, col, side) for side in ("left", "right", "top", "bottom"))

    def own_edges(self, row, col):
        style = self.style(row, col)
        return sum(bool(style.get(key)) for key in ("bl", "br", "bt", "bb"))


def detect_fields(layout):
    """A first guess at which cells are filled in, and with what.

    * Empty boxed cells in columns that are otherwise blank are the readings.
      A column that already has text in a fair share of its boxes -- Sr No,
      Parameters, UOM -- is a label column, and its blanks stay blank.
    * A reading's type comes from its row: a unit means a number, 'Absent' or
      'Present / Absent' means a choice, a spec like '6.5 - 8.5' sets limits.
      A reading row with no label of its own (the row under a 'Time' header)
      takes its type from the header above it instead.
    * Labels like 'Date:', 'Remarks:', 'Q.A Chemist' and 'Q.A.M' bind the cell
      beside them (or below, at the right edge) to the record itself.

    Every guess is only a starting point; the manager corrects it on screen.
    """
    grid = _Grid(layout)
    fields = {}

    # -- label columns ----------------------------------------------------
    boxed = {}
    for row in range(grid.min_row, grid.max_row + 1):
        for col in range(grid.min_col, grid.max_col + 1):
            if grid.covered(row, col):
                continue
            if grid.edges(row, col) >= 3:
                boxed[(row, col)] = grid.text(row, col) != ""

    label_columns = set()
    for col in range(grid.min_col, grid.max_col + 1):
        column = [filled for (row, c), filled in boxed.items() if c == col]
        if column and sum(column) / len(column) > 0.3:
            label_columns.add(col)

    # -- the table's header row, to know which column is name / unit / spec --
    header_row = None
    for row in range(grid.min_row, grid.max_row + 1):
        labels = [col for (r, col), filled in boxed.items() if r == row and filled]
        if len(labels) >= 3:
            header_row = row
            break
    name_col = unit_col = spec_col = None
    if header_row is not None:
        for col in range(grid.min_col, grid.max_col + 1):
            heading = grid.text(header_row, col, follow_merge=False)
            if not heading:
                continue
            if unit_col is None and _UNIT_HEADER.match(heading):
                unit_col = col
            elif spec_col is None and _SPEC_HEADER.search(heading):
                spec_col = col
            elif name_col is None and _NAME_HEADER.search(heading):
                name_col = col

    def row_label(row):
        """The row's own labels (not text spilling in from a merge above)."""
        texts = {
            col: grid.text(row, col, follow_merge=False)
            for col in sorted(label_columns)
            if grid.text(row, col, follow_merge=False)
        }
        if not texts:
            return None
        if name_col in texts:
            name = texts[name_col]
        else:
            words = [text for text in texts.values() if re.search(r"[A-Za-z]", text)]
            name = max(words, key=len) if words else ""
        return {
            "name": re.sub(r"\s+", " ", name).strip(),
            "unit": texts.get(unit_col, "") if unit_col else "",
            "spec": texts.get(spec_col, "") if spec_col else "",
        }

    def column_heading(row, col):
        for above in range(row - 1, grid.min_row - 1, -1):
            text = grid.text(above, col)
            if text:
                return re.sub(r"\s+", " ", text).strip()
        return ""

    # -- readings -------------------------------------------------------
    readings = []
    for (row, col), filled in sorted(boxed.items()):
        if filled or col in label_columns:
            continue
        readings.append((row, col))

    heading_counts = {}
    for row, col in readings:
        field = {"type": FieldType.TEXT}
        label = row_label(row)
        heading = column_heading(row, col)
        directly_above = grid.text(row - 1, col) if grid.inside(row - 1, col) else ""

        if label is None:
            if _TIME_LABEL.match(directly_above):
                field["type"] = FieldType.TIME
            elif _DATE_LABEL.match(directly_above):
                field["type"] = FieldType.DATE
        else:
            choice = _choice_from(label["spec"]) or _choice_from(label["unit"])
            low, high = parse_specification(label["spec"])
            if choice:
                field["type"] = FieldType.CHOICE
                field["options"], field["ok"] = choice
            elif re.search(r"\btime\b", label["name"], re.I):
                field["type"] = FieldType.TIME
            elif re.search(r"\bdate\b", label["name"], re.I):
                field["type"] = FieldType.DATE
            elif low is not None or high is not None or (label["unit"] and len(label["unit"]) <= 15):
                field["type"] = FieldType.NUMBER
            if low is not None:
                field["min"] = low
            if high is not None:
                field["max"] = high
            if label["spec"]:
                field["spec"] = label["spec"][:160]

        # "Free Fatty Acids · Time 3": the row, then which column.
        ordinal_key = (row, heading)
        heading_counts[ordinal_key] = heading_counts.get(ordinal_key, 0) + 1
        column_part = f"{heading} {heading_counts[ordinal_key]}".strip() if heading else ""
        parts = [label["name"]] if label and label["name"] else []
        if column_part:
            parts.append(column_part)
        if parts:
            field["label"] = " · ".join(parts)[:200]
        fields[ref(row, col)] = field

    # -- labels that bind a neighbouring cell to the record ---------------
    def free(row, col):
        return (
            grid.inside(row, col)
            and not grid.covered(row, col)
            and grid.text(row, col) == ""
            and ref(row, col) not in fields
        )

    bindings = []
    for address, entry in grid.cells.items():
        text = (entry.get("v") or "").strip()
        if not text or len(text) > 40:
            continue
        row, col = parse_ref(address)
        # A heading inside a table ("Date" over a column of dates) is not a
        # label for the record; only loose labels or 'Label:' forms are.
        if grid.own_edges(row, col) >= 2 and not text.endswith(":"):
            continue
        if _DATE_LABEL.match(text):
            kind = FieldType.RECORD_DATE
        elif _SHIFT_LABEL.match(text):
            kind = FieldType.SHIFT
        elif _REMARKS_LABEL.match(text):
            kind = FieldType.REMARKS
        elif _SUBMITTER_LABEL.search(text):
            kind = FieldType.SIGN_SUBMITTED
        elif _APPROVER_LABEL.search(text):
            kind = FieldType.SIGN_APPROVED
        else:
            continue
        bindings.append((row, col, kind, text))

    remarks_taken = False
    for row, col, kind, text in sorted(bindings):
        if kind == FieldType.REMARKS and remarks_taken:
            continue
        _, c2 = grid.span(row, col)
        r2, _ = grid.span(row, col)
        target = None
        if free(row, c2 + 1):
            target = (row, c2 + 1)
        elif free(r2 + 1, col):
            target = (r2 + 1, col)
        if target is None:
            continue
        fields[ref(*target)] = {"type": kind, "label": text.rstrip(":").strip()[:200]}
        remarks_taken = remarks_taken or kind == FieldType.REMARKS

    return fields


# ---------------------------------------------------------------------------
# Header guess (document code, revision, title...)
# ---------------------------------------------------------------------------

_CODE_RE = re.compile(r"\b([A-Z]{2,}(?:-[A-Z0-9]+){2,})\b")
_REVISION_RE = re.compile(
    r"rev(?:ision)?\.?\s*(?:no\.?)?\s*[:\-]?\s*(\d{1,3})"
    r"(?:\s*[/\-]\s*(\d{1,2})[-./](\d{1,2})[-./](\d{2,4}))?",
    re.I,
)
_CLASSIFICATION_RE = re.compile(r"classifi(?:ed|cation)\s*:?\s*([^\n]+)", re.I)


def _clean(text):
    return re.sub(r"\s+", " ", text or "").strip()


def guess_header(layout, filename=""):
    header_texts = [
        text
        for part in (layout.get("header", {}), layout.get("footer", {}))
        for text in part.values()
        if text
    ]
    cell_texts = [entry["v"] for entry in layout["cells"].values() if entry.get("v")]
    stem = re.sub(r"\.(xlsx|xlsm)$", "", filename or "", flags=re.I)

    document_code = ""
    for text in header_texts + cell_texts + [stem]:
        match = _CODE_RE.search(text.upper() if text is stem else text)
        if match and re.search(r"\d", match.group(1)):
            document_code = match.group(1)
            break

    revision_number, revision_date = "", None
    for text in header_texts + cell_texts:
        match = _REVISION_RE.search(text)
        if match:
            revision_number = match.group(1).zfill(2)
            if match.group(2):
                day, month, year = int(match.group(2)), int(match.group(3)), int(match.group(4))
                year = year + 2000 if year < 100 else year
                try:
                    revision_date = date(year, month, day).isoformat()
                except ValueError:
                    revision_date = None
            break

    classification = ""
    for text in header_texts + cell_texts:
        match = _CLASSIFICATION_RE.search(text)
        if match:
            classification = _clean(match.group(1))[:120]
            break

    def is_meta(text):
        return bool(
            _CODE_RE.search(text)
            or _REVISION_RE.search(text)
            or _CLASSIFICATION_RE.search(text)
            or "controlled document" in text.lower()
        )

    organisation = ""
    header = layout.get("header", {})
    for position in ("center", "left", "right"):
        text = _clean(header.get(position, ""))
        if text and not is_meta(text):
            organisation = text[:255]
            break

    # The title is the biggest text above the table that is not a label. Rows
    # inside the table are left out: their labels are often set larger than
    # the title itself.
    grid = _Grid(layout)
    last_title_row = grid.min_row + 5
    for row in range(grid.min_row, grid.max_row + 1):
        boxed_labels = sum(
            1
            for col in range(grid.min_col, grid.max_col + 1)
            if not grid.covered(row, col)
            and grid.text(row, col)
            and grid.own_edges(row, col) >= 3
        )
        if boxed_labels >= 3:
            if row > grid.min_row:
                last_title_row = row - 1
            break
    best = None
    for address, entry in layout["cells"].items():
        text = _clean(entry.get("v", ""))
        row, _ = parse_ref(address)
        if not text or row > last_title_row or len(text) < 6 or is_meta(text):
            continue
        if text.endswith(":") or text == organisation:
            continue
        size = layout["styles"][entry["s"]].get("fs", 11)
        width = 1
        if address in grid.merge_of:
            r1, c1, r2, c2 = grid.merge_of[address]
            width = c2 - c1 + 1
        rank = (size, width, -row)
        if best is None or rank > best[0]:
            best = (rank, text)
    title = best[1] if best else ""
    if not title and stem:
        title = _clean(_CODE_RE.sub("", stem.upper()))

    return {
        "document_code": document_code,
        "title": title[:255],
        "organisation": organisation,
        "revision_number": revision_number,
        "revision_date": revision_date,
        "classification": classification,
    }


# ---------------------------------------------------------------------------
# Reading an upload
# ---------------------------------------------------------------------------


def read_sheet(uploaded, filename="", sheet_name=None):
    """Parse an uploaded workbook into everything the designer needs."""
    name = filename or getattr(uploaded, "name", "") or ""
    if name.lower().endswith(".xls"):
        raise SheetImportError(
            "This is an old-format .xls file. Open it in Excel, save it as "
            "'Excel Workbook (.xlsx)', and upload that."
        )
    try:
        workbook = load_workbook(uploaded, data_only=True)
    except (InvalidFileException, KeyError, OSError, ValueError, TypeError) as error:
        raise SheetImportError(
            "This file could not be read as an Excel workbook (.xlsx)."
        ) from error
    except Exception as error:  # zipfile.BadZipFile and friends
        raise SheetImportError(
            "This file could not be read as an Excel workbook (.xlsx)."
        ) from error

    sheets = [ws.title for ws in workbook.worksheets]
    if sheet_name:
        if sheet_name not in sheets:
            raise SheetImportError(f"The workbook has no sheet named '{sheet_name}'.")
        worksheet = workbook[sheet_name]
    else:
        worksheet = workbook.active if workbook.active in workbook.worksheets else workbook.worksheets[0]

    layout = build_layout(worksheet, _theme_palette(workbook))
    return {
        "sheets": sheets,
        "sheet": worksheet.title,
        "layout": layout,
        "layout_token": sign_layout(layout),
        "cell_fields": detect_fields(layout),
        "header": guess_header(layout, name),
    }


# ---------------------------------------------------------------------------
# Saving: validating fields and values
# ---------------------------------------------------------------------------


def _decimal_or_none(value, label):
    if value in (None, ""):
        return None
    try:
        return str(Decimal(str(value).strip()))
    except (InvalidOperation, ValueError):
        raise ValueError(f"{label} must be a number.")


def _string_list(value, label):
    if value in (None, ""):
        return []
    if not isinstance(value, list) or len(value) > 50:
        raise ValueError(f"{label} must be a list of up to 50 options.")
    cleaned = []
    for item in value:
        text = str(item).strip()[:100]
        if text and text not in cleaned:
            cleaned.append(text)
    return cleaned


def clean_cell_fields(fields, layout):
    """Validate the manager's cell fields against the layout.

    Returns ``(cleaned, errors)``; ``errors`` maps a cell (or '__all__') to a
    message. Unknown keys are dropped rather than stored.
    """
    if not isinstance(fields, dict):
        return {}, {"__all__": "Cell fields must be an object keyed by cell."}
    if len(fields) > 5000:
        return {}, {"__all__": "Too many fillable cells."}
    grid = _Grid(layout)
    cleaned, errors = {}, {}
    remarks = []
    for address, spec in fields.items():
        try:
            row, col = parse_ref(address)
        except ValueError as error:
            errors[str(address)[:20]] = str(error)
            continue
        if not grid.inside(row, col):
            errors[address] = "is outside the sheet."
            continue
        if grid.covered(row, col):
            errors[address] = f"is inside the merged cell {grid.anchor_of[address]}."
            continue
        if not isinstance(spec, dict) or spec.get("type") not in FIELD_TYPES:
            errors[address] = "has no valid type."
            continue
        kind = spec["type"]
        entry = {"type": kind}
        label = str(spec.get("label") or "").strip()[:200]
        if label:
            entry["label"] = label
        try:
            if kind == FieldType.NUMBER:
                low = _decimal_or_none(spec.get("min"), "Min")
                high = _decimal_or_none(spec.get("max"), "Max")
                if low is not None and high is not None and Decimal(low) > Decimal(high):
                    raise ValueError("Min is above max.")
                if low is not None:
                    entry["min"] = low
                if high is not None:
                    entry["max"] = high
            if kind == FieldType.CHOICE:
                options = _string_list(spec.get("options"), "Options")
                ok = _string_list(spec.get("ok"), "Conforming values")
                if options:
                    entry["options"] = options
                if ok:
                    entry["ok"] = ok
        except ValueError as error:
            errors[address] = str(error)
            continue
        spec_text = str(spec.get("spec") or "").strip()[:160]
        if spec_text and kind in VALUE_TYPES:
            entry["spec"] = spec_text
        if kind == FieldType.REMARKS:
            remarks.append(address)
        cleaned[address] = entry
    if len(remarks) > 1:
        errors["__all__"] = (
            f"Only one cell can hold the remarks; {', '.join(sorted(remarks))} all do."
        )
    return cleaned, errors


def clean_cell_value(field, raw):
    """Normalise one typed value, or raise ValueError saying what is wrong."""
    value = "" if raw is None else str(raw).strip()
    if len(value) > MAX_VALUE_LENGTH:
        raise ValueError(f"is longer than {MAX_VALUE_LENGTH} characters.")
    if value == "":
        return ""
    kind = field.get("type")
    if kind == FieldType.TIME:
        match = re.fullmatch(r"(\d{1,2})[:.](\d{2})(?::\d{2})?", value)
        if not match or int(match.group(1)) > 23 or int(match.group(2)) > 59:
            raise ValueError("is not a time (HH:MM).")
        return f"{int(match.group(1)):02d}:{match.group(2)}"
    if kind == FieldType.DATE:
        try:
            return date.fromisoformat(value).isoformat()
        except ValueError:
            raise ValueError("is not a date (YYYY-MM-DD).")
    return value


def check_cell(field, raw):
    """True (meets spec), False (does not), None (blank or not judged).

    Mirrors ``RecordTemplateParameter.check_value``: a number is range-checked
    only when limits are set, and a choice is judged against the conforming
    values, never against the offered options (those include the failing one).
    """
    if raw is None or str(raw).strip() == "":
        return None
    value = str(raw).strip()
    kind = field.get("type")
    if kind == FieldType.CHOICE:
        ok = field.get("ok") or []
        if not ok:
            return None
        return any(value.casefold() == str(option).casefold() for option in ok)
    if kind == FieldType.NUMBER:
        low, high = field.get("min"), field.get("max")
        if low in (None, "") and high in (None, ""):
            return None
        try:
            number = Decimal(value)
        except (InvalidOperation, ValueError):
            return False
        if low not in (None, "") and number < Decimal(str(low)):
            return False
        if high not in (None, "") and number > Decimal(str(high)):
            return False
        return True
    return None
