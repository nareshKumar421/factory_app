"""Shift sheet: the floor's Excel read into rows, planned, and booked as runs."""
import io
from datetime import date, datetime, timedelta
from decimal import Decimal
from types import SimpleNamespace
from zoneinfo import ZoneInfo

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Permission
from django.test import SimpleTestCase, TestCase
from openpyxl import Workbook
from rest_framework.test import APIClient

from company.models import Company, UserCompany, UserRole

from .models import BlowingMachine, BlowingRun, BlowingSegment, PreformSpec
from .services import shift_sheet as ss

IST = ZoneInfo('Asia/Kolkata')

HEADER = ['sku', 'shift', 'total production', 'labour', None, 'total electricity', 'utility',
          'wastage']
SUBHEADER = [None, None, None, 'company', 'outside', None, None, None]


def _xlsx(lines, *, start_row=1):
    """A workbook holding ``lines`` from ``start_row``, the way the floor lays it out."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Sheet1'
    for offset, line in enumerate(lines):
        for column, value in enumerate(line, start=1):
            if value is not None:
                sheet.cell(row=start_row + offset, column=column, value=value)
    workbook.create_sheet('Sheet2')
    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    buffer.name = 'sidel data.xlsx'
    return buffer


def _the_floor_sheet():
    """``sidel data.xlsx`` of 30.09.2026 as it came in: rows 8–12, and a stray
    meter digit on row 41."""
    lines = [
        ['30.09.2026'],
        HEADER,
        SUBHEADER,
        ['frystal 40 gms'],
        [None, 'night', 27440, 0, 5, 146.4, 765.6, 20],
    ] + [[]] * 28 + [
        [None, None, None, None, None, 1, None, None],
    ]
    return _xlsx(lines, start_row=8)


class _Fixture(TestCase):
    def setUp(self):
        self.company = Company.objects.create(name='Test Oil', code='TEST_OIL_SS')
        self.machine = BlowingMachine.objects.create(company=self.company, name='Synergy 1/4')
        self.f40 = PreformSpec.objects.create(
            company=self.company, make='Frystal', gram=Decimal('40'), preforms_per_box=600,
            preform_rate_per_bottle=Decimal('6.18'), sap_item_code='PM0000594')
        self.f21 = PreformSpec.objects.create(
            company=self.company, make='Frystal', gram=Decimal('21'), preforms_per_box=1152,
            preform_rate_per_bottle=Decimal('3.2445'))
        self.p49 = PreformSpec.objects.create(
            company=self.company, make='Pioneer', gram=Decimal('49.5'), preforms_per_box=490)

    def _run(self, day, number, *, started=None, ended=None, start=None, stop=None,
             spec=None, production=20000, remarks=''):
        run = BlowingRun.objects.create(
            company=self.company, machine=self.machine, preform_spec=spec or self.f40,
            date=day, run_number=number, status='COMPLETED', remarks=remarks,
            machine_start_reading=start, machine_stop_reading=stop,
            total_counter_production=production,
        )
        if started:
            BlowingSegment.objects.create(
                blowing_run=run, start_time=started, end_time=ended, is_active=False)
        return run

    def _row(self, day, shift, **extra):
        row = {
            'date': day.isoformat(), 'shift': shift, 'preform_spec_id': self.f40.id,
            'total_counter_production': '27440', 'own_labour_count': '0',
            'contract_labour_count': '5', 'machine_units': '146.4',
            'utility_units': '765.6', 'rejection_pcs': '20',
        }
        row.update(extra)
        return row

    def _plan(self, *rows):
        return ss.plan_rows(self.company, self.machine, list(rows))


def _ist(y, m, d, hh, mm):
    return datetime(y, m, d, hh, mm, tzinfo=IST)


# ===========================================================================
# Reading the sheet
# ===========================================================================

class ParseWorkbookTest(_Fixture):

    def test_reads_the_floor_sheet_as_it_came_in(self):
        parsed = ss.parse_workbook(_the_floor_sheet(), self.company)
        self.assertEqual(parsed['rows'], [{
            'sheet': 'Sheet1', 'line': 12, 'date': '2026-09-30', 'shift': 'NIGHT',
            'sku_text': 'frystal 40 gms', 'preform_spec_id': self.f40.id,
            'total_counter_production': '27440', 'own_labour_count': '0',
            'contract_labour_count': '5', 'machine_units': '146.4',
            'utility_units': '765.6', 'rejection_pcs': '20', 'notes': [],
        }])
        # The stray digit under total electricity is not guessed into a run.
        self.assertEqual(len(parsed['ignored']), 1)
        self.assertEqual(parsed['ignored'][0]['line'], 41)
        self.assertIn('no shift', parsed['ignored'][0]['reason'])

    def test_several_days_and_skus_in_one_sheet(self):
        parsed = ss.parse_workbook(_xlsx([
            ['28.09.2026'], HEADER, SUBHEADER,
            ['frystal 21 gms'],
            [None, 'day', 10475, 0, 6, 51.67, 409.32, 10],
            [None, 'night', 9677, 0, 6, 70, 731.2, 35],
            ['pioneer 49.5 gm'],
            [None, 'night', 500, 0, 1, 5, 10, 0],
            [datetime(2026, 9, 29)],
            ['frystal 40 gms'],
            [None, 'Day Shift', 21032, 1, 5, 121.65, 756.8, 10],
        ]), self.company)
        got = [(r['date'], r['shift'], r['preform_spec_id'], r['total_counter_production'])
               for r in parsed['rows']]
        self.assertEqual(got, [
            ('2026-09-28', 'DAY', self.f21.id, '10475'),
            ('2026-09-28', 'NIGHT', self.f21.id, '9677'),
            ('2026-09-28', 'NIGHT', self.p49.id, '500'),
            ('2026-09-29', 'DAY', self.f40.id, '21032'),
        ])
        self.assertEqual(parsed['rows'][-1]['own_labour_count'], '1')
        self.assertEqual(parsed['ignored'], [])

    def test_what_it_cannot_match_is_left_for_the_person(self):
        parsed = ss.parse_workbook(_xlsx([
            HEADER, SUBHEADER,
            ['mystery 33 gms'],
            [None, 'evening', 100, 0, 1, 1, 1, 'some'],
        ]), self.company)
        row = parsed['rows'][0]
        self.assertIsNone(row['date'])
        self.assertIsNone(row['shift'])
        self.assertIsNone(row['preform_spec_id'])
        self.assertIsNone(row['rejection_pcs'])
        notes = ' '.join(row['notes'])
        self.assertIn('neither day nor night', notes)
        self.assertIn('No date', notes)
        self.assertIn("matches no preform spec", notes)
        self.assertIn("'some' is not a number", notes)

    def test_a_shift_with_no_figures_is_left_out(self):
        parsed = ss.parse_workbook(_xlsx([
            ['30.09.2026'], HEADER, SUBHEADER,
            ['frystal 40 gms'],
            [None, 'day'],
            [None, 'night', 27440, 0, 5, 146.4, 765.6, 20],
        ]), self.company)
        self.assertEqual([r['shift'] for r in parsed['rows']], ['NIGHT'])
        self.assertIn('Day shift with no figures', parsed['ignored'][0]['reason'])

    def test_a_labour_column_without_the_split_is_outside_labour(self):
        parsed = ss.parse_workbook(_xlsx([
            ['30.09.2026'],
            ['sku', 'shift', 'total production', 'labour', 'total electricity', 'utility',
             'wastage'],
            ['frystal 40 gms', 'night', 27440, 5, 146.4, 765.6, 20],
        ]), self.company)
        row = parsed['rows'][0]
        self.assertEqual(row['contract_labour_count'], '5')
        self.assertIsNone(row['own_labour_count'])
        self.assertEqual(row['machine_units'], '146.4')

    def test_not_a_workbook(self):
        junk = io.BytesIO(b'not an excel file')
        with self.assertRaisesMessage(ss.SheetError, 'could not be opened'):
            ss.parse_workbook(junk, self.company)

    def test_a_workbook_with_no_shift_rows(self):
        with self.assertRaisesMessage(ss.SheetError, 'No shift rows found'):
            ss.parse_workbook(_xlsx([['hello'], ['world']]), self.company)


class MatchSpecTest(SimpleTestCase):
    SPECS = [
        SimpleNamespace(id=7, make='Frystal', gram=Decimal('40.00')),
        SimpleNamespace(id=9, make='Pioneer', gram=Decimal('49.50')),
        SimpleNamespace(id=10, make='pioneer', gram=Decimal('40.00')),
        SimpleNamespace(id=11, make='Frystal', gram=Decimal('21.00')),
    ]

    def match(self, text):
        spec, note = ss.match_spec(text, self.SPECS)
        return (spec.id if spec else None), note

    def test_make_and_gram(self):
        self.assertEqual(self.match('frystal 40 gms'), (7, None))
        self.assertEqual(self.match('Pioneer 40 gms'), (10, None))
        self.assertEqual(self.match('pioneer 49.5 gm'), (9, None))
        self.assertEqual(self.match('FRYSTAL 21g'), (11, None))

    def test_no_or_two_matches_are_not_guessed(self):
        spec, note = self.match('frystal 50 gms')
        self.assertIsNone(spec)
        self.assertIn('matches no preform spec', note)
        twins = self.SPECS + [SimpleNamespace(id=12, make='Frystal', gram=Decimal('40'))]
        spec, note = ss.match_spec('frystal 40', twins)
        self.assertIsNone(spec)
        self.assertIn('matches 2 preform specs', note)


# ===========================================================================
# Which shift a run in the app ran
# ===========================================================================

class RunShiftTest(_Fixture):

    def test_read_off_the_first_segment_not_the_run_number(self):
        # 29 Sep: the floor's run 1 that ran 20:35 → 07:44 is the night shift.
        night = self._run(date(2026, 9, 29), 1, started=_ist(2026, 9, 29, 20, 35),
                          ended=_ist(2026, 9, 30, 7, 44))
        day = self._run(date(2026, 9, 30), 1, started=_ist(2026, 9, 30, 7, 56),
                        ended=_ist(2026, 9, 30, 19, 7))
        self.assertEqual(ss.run_shift(night), ss.NIGHT)
        self.assertEqual(ss.run_shift(day), ss.DAY)

    def test_a_sheet_booked_run_says_so_in_its_remarks(self):
        run = self._run(date(2026, 9, 11), 1, started=_ist(2026, 9, 11, 17, 54),
                        remarks='Backfilled from sidel data.xlsx — night shift')
        self.assertEqual(ss.run_shift(run), ss.NIGHT)

    def test_with_no_segment_it_is_when_the_run_was_opened(self):
        run = self._run(date(2026, 9, 30), 1)
        BlowingRun.objects.filter(pk=run.pk).update(created_at=_ist(2026, 9, 30, 7, 44))
        run.refresh_from_db()
        self.assertEqual(ss.run_shift(run), ss.DAY)


# ===========================================================================
# Planning
# ===========================================================================

class PlanRowsTest(_Fixture):

    def setUp(self):
        super().setUp()
        self.day30 = self._run(
            date(2026, 9, 30), 1, started=_ist(2026, 9, 30, 7, 56),
            ended=_ist(2026, 9, 30, 19, 7),
            start=Decimal('34682.1824'), stop=Decimal('34812.3424'), production=22497)

    def test_the_night_row_carries_on_from_the_day_run(self):
        plan = self._plan(self._row(date(2026, 9, 30), 'NIGHT'))
        row = plan['rows'][0]
        self.assertEqual(row['status'], 'NEW', row)
        self.assertEqual(row['run_number'], 2)
        self.assertEqual(row['machine_start_reading'], '34812.3424')
        self.assertEqual(row['machine_stop_reading'], '34958.7424')
        self.assertEqual(row['good_bottles'], 27420)
        self.assertEqual(plan['summary'], {'NEW': 1, 'DUPLICATE': 0, 'ERROR': 0})
        self.assertEqual([e['id'] for e in plan['existing']], [self.day30.id])
        self.assertEqual(plan['existing'][0]['shift'], 'DAY')

    def test_a_shift_the_floor_entered_is_skipped(self):
        # Run 1 of the 29th is that day's night shift — a run-number check
        # would have let the sheet's night row in beside it.
        night29 = self._run(
            date(2026, 9, 29), 1, started=_ist(2026, 9, 29, 20, 35),
            ended=_ist(2026, 9, 30, 7, 44),
            start=Decimal('34560.3520'), stop=Decimal('34682.0032'))
        row = self._plan(self._row(date(2026, 9, 29), 'NIGHT'))['rows'][0]
        self.assertEqual(row['status'], 'DUPLICATE')
        self.assertEqual(row['duplicate_of']['id'], night29.id)
        self.assertIsNone(row['run_number'])

    def test_add_anyway_books_it_beside_the_floor_run(self):
        self._run(date(2026, 9, 29), 1, started=_ist(2026, 9, 29, 20, 35))
        row = self._plan(self._row(date(2026, 9, 29), 'NIGHT', add_anyway=True))['rows'][0]
        self.assertEqual(row['status'], 'NEW')
        self.assertEqual(row['run_number'], 2)
        self.assertIn('beside run 1', row['warnings'][0])

    def test_another_sku_in_the_same_shift_is_booked_with_a_warning(self):
        self._run(date(2026, 9, 29), 1, started=_ist(2026, 9, 29, 20, 35))
        row = self._plan(
            self._row(date(2026, 9, 29), 'NIGHT', preform_spec_id=self.f21.id))['rows'][0]
        self.assertEqual(row['status'], 'NEW')
        self.assertIn('already has run 1 for this shift', row['warnings'][0])

    def test_rows_chain_in_shift_order_whatever_order_they_are_typed(self):
        plan = self._plan(
            self._row(date(2026, 9, 30), 'NIGHT', machine_units='100'),
            self._row(date(2026, 9, 29), 'DAY', machine_units='10'),
        )
        night30, day29 = plan['rows']
        # 29 Sep day sits before the 30th's day run: it starts from nothing
        # earlier on the machine, so the meter starts at 0 and says so.
        self.assertEqual(day29['machine_start_reading'], '0')
        self.assertIn('No earlier meter reading', day29['warnings'][0])
        self.assertIn('starts the meter at 34682.1824', day29['warnings'][-1])
        self.assertEqual(day29['run_number'], 1)
        self.assertEqual(night30['machine_start_reading'], '34812.3424')
        self.assertEqual(night30['machine_stop_reading'], '34912.3424')

    def test_a_draft_with_no_readings_does_not_hide_the_last_reading(self):
        # The last day before the sheet holds only a draft that was never
        # finished; the reading to carry on from is the day before that.
        self._run(date(2026, 9, 20), 1, start=Decimal('900'), stop=Decimal('1000'))
        self._run(date(2026, 9, 25), 1)
        row = self._plan(self._row(date(2026, 9, 27), 'DAY'))['rows'][0]
        self.assertEqual(row['machine_start_reading'], '1000.0000')
        self.assertFalse(any('No earlier meter reading' in w for w in row['warnings']))

    def test_a_row_slotted_before_a_later_run_checks_the_meter_meets(self):
        later = self._run(date(2026, 10, 1), 1, started=_ist(2026, 10, 1, 8, 0),
                          start=Decimal('34958.7424'), stop=Decimal('35000'))
        row = self._plan(self._row(date(2026, 9, 30), 'NIGHT'))['rows'][0]
        self.assertEqual(row['warnings'], [])
        BlowingRun.objects.filter(pk=later.pk).update(machine_start_reading=Decimal('34990'))
        row = self._plan(self._row(date(2026, 9, 30), 'NIGHT'))['rows'][0]
        self.assertIn('starts the meter at 34990', row['warnings'][0])

    def test_what_must_be_fixed_first(self):
        tomorrow = date.today() + timedelta(days=2)
        plan = self._plan(
            self._row(date(2026, 9, 30), None),
            self._row(date(2026, 9, 30), 'NIGHT', preform_spec_id=None),
            self._row(tomorrow, 'DAY'),
            self._row(date(2026, 9, 30), 'NIGHT', rejection_pcs='30000'),
            self._row(date(2026, 9, 30), 'NIGHT', total_counter_production='0'),
            self._row(date(2026, 9, 30), 'NIGHT', machine_units='abc'),
        )
        errors = [' '.join(r['errors']) for r in plan['rows']]
        self.assertTrue(all(r['status'] == 'ERROR' for r in plan['rows']))
        self.assertIn('Pick the shift', errors[0])
        self.assertIn('Pick the preform', errors[1])
        self.assertIn('is in the future', errors[2])
        self.assertIn('more than the total production', errors[3])
        self.assertIn('more than zero', errors[4])
        self.assertIn("'abc' is not a number", errors[5])

    def test_the_same_shift_twice_in_one_sheet(self):
        plan = self._plan(self._row(date(2026, 9, 30), 'NIGHT'),
                          self._row(date(2026, 9, 30), 'NIGHT'))
        self.assertEqual([r['status'] for r in plan['rows']], ['NEW', 'ERROR'])
        self.assertIn('Row 1 already has', plan['rows'][1]['errors'][0])

    def test_planning_writes_nothing(self):
        self._plan(self._row(date(2026, 9, 30), 'NIGHT'))
        self.assertEqual(BlowingRun.objects.filter(company=self.company).count(), 1)


# ===========================================================================
# Booking
# ===========================================================================

class ApplyRowsTest(_Fixture):

    def setUp(self):
        super().setUp()
        self.user = get_user_model().objects.create_user(email='arvin@test.com', password='pw')
        self._run(date(2026, 9, 30), 1, started=_ist(2026, 9, 30, 7, 56),
                  start=Decimal('34682.1824'), stop=Decimal('34812.3424'))

    def test_books_completed_runs_dated_to_their_shift(self):
        result = ss.apply_rows(
            self.company, self.machine.id,
            [self._row(date(2026, 9, 30), 'NIGHT'), self._row(date(2026, 10, 1), 'DAY',
                                                              machine_units='10')],
            self.user, source='sidel data.xlsx')
        self.assertEqual([(c['date'], c['shift'], c['run_number']) for c in result['created']],
                         [('2026-09-30', 'NIGHT', 2), ('2026-10-01', 'DAY', 1)])
        night = BlowingRun.objects.get(id=result['created'][0]['id'])
        self.assertEqual(night.status, 'COMPLETED')
        self.assertEqual(night.warehouse_approval_status, 'APPROVED')
        self.assertEqual(night.created_at, _ist(2026, 9, 30, 20, 0))
        self.assertEqual(night.created_by, self.user)
        self.assertEqual(night.remarks,
                         'Entered from the shift sheet (sidel data.xlsx) — night shift')
        self.assertEqual(night.machine_start_reading, Decimal('34812.3424'))
        self.assertEqual(night.machine_stop_reading, Decimal('34958.7424'))
        self.assertEqual(night.machine_units, Decimal('146.4'))
        self.assertEqual(night.utility_units, Decimal('765.6'))
        self.assertEqual((night.operator_count, night.own_labour_count,
                          night.contract_labour_count), (1, 0, 5))
        self.assertEqual(night.rejection_pcs, 20)
        self.assertEqual(night.preform_rate_per_bottle, Decimal('6.18'))
        self.assertEqual(night.sap_preform_item_code, 'PM0000594')
        self.assertEqual(night.cost_summary.good_bottles, 27420)
        day = BlowingRun.objects.get(id=result['created'][1]['id'])
        self.assertEqual(day.created_at, _ist(2026, 10, 1, 9, 0))
        self.assertEqual(day.machine_start_reading, Decimal('34958.7424'))

        # The same sheet a second time adds nothing.
        again = self._plan(self._row(date(2026, 9, 30), 'NIGHT'))
        self.assertEqual(again['rows'][0]['status'], 'DUPLICATE')

    def test_one_bad_row_books_nothing(self):
        with self.assertRaises(ss.ShiftSheetRefused) as caught:
            ss.apply_rows(self.company, self.machine.id, [
                self._row(date(2026, 9, 30), 'NIGHT'),
                self._row(date(2026, 9, 30), 'DAY', preform_spec_id=None),
            ], self.user)
        self.assertEqual(caught.exception.plan['summary']['ERROR'], 1)
        self.assertNotIn('_cleaned', caught.exception.plan)
        self.assertEqual(BlowingRun.objects.filter(company=self.company).count(), 1)

    def test_nothing_new_is_refused(self):
        self._run(date(2026, 9, 29), 1, started=_ist(2026, 9, 29, 20, 35))
        with self.assertRaisesMessage(ss.ShiftSheetRefused, 'Nothing to add'):
            ss.apply_rows(self.company, self.machine.id,
                          [self._row(date(2026, 9, 29), 'NIGHT')], self.user)


# ===========================================================================
# API
# ===========================================================================

class ShiftSheetAPITest(_Fixture):
    URL = '/api/v1/blowing/shift-sheet/'
    PARSE_URL = '/api/v1/blowing/shift-sheet/parse/'

    def setUp(self):
        super().setUp()
        self.user = get_user_model().objects.create_user(email='entry@test.com', password='pw')
        UserCompany.objects.create(
            user=self.user, company=self.company,
            role=UserRole.objects.create(name='Staff'), is_active=True)
        self.client = APIClient()
        self.client.credentials(HTTP_COMPANY_CODE=self.company.code)
        self.client.force_authenticate(user=self.user)

    def grant(self, *codenames):
        self.user.user_permissions.set(Permission.objects.filter(
            content_type__app_label='blowing', codename__in=codenames))
        self.user = get_user_model().objects.get(pk=self.user.pk)
        self.client.force_authenticate(user=self.user)

    def body(self, **extra):
        return {'machine_id': self.machine.id, 'rows': [self._row(date(2026, 9, 30), 'NIGHT')],
                **extra}

    def test_needs_create_and_complete(self):
        self.grant('can_create_blowing_run')
        self.assertEqual(self.client.post(self.URL, self.body(), format='json').status_code, 403)
        self.assertEqual(
            self.client.post(self.PARSE_URL, {'file': _the_floor_sheet()},
                             format='multipart').status_code, 403)

    def test_upload_check_then_save(self):
        self.grant('can_create_blowing_run', 'can_complete_blowing_run')
        parsed = self.client.post(self.PARSE_URL, {'file': _the_floor_sheet()},
                                  format='multipart')
        self.assertEqual(parsed.status_code, 200, parsed.data)
        self.assertEqual(parsed.data['file_name'], 'sidel data.xlsx')
        rows = parsed.data['rows']
        self.assertEqual(len(rows), 1)

        checked = self.client.post(
            self.URL, {'machine_id': self.machine.id, 'rows': rows}, format='json')
        self.assertEqual(checked.status_code, 200, checked.data)
        self.assertFalse(checked.data['committed'])
        self.assertEqual(checked.data['rows'][0]['status'], 'NEW')
        self.assertFalse(BlowingRun.objects.filter(company=self.company).exists())

        saved = self.client.post(self.URL, {
            'machine_id': self.machine.id, 'rows': rows, 'commit': True,
            'source': parsed.data['file_name'],
        }, format='json')
        self.assertEqual(saved.status_code, 201, saved.data)
        self.assertTrue(saved.data['committed'])
        run = BlowingRun.objects.get(id=saved.data['created'][0]['id'])
        self.assertEqual(run.created_by, self.user)
        self.assertIn('(sidel data.xlsx) — night shift', run.remarks)

    def test_a_row_to_fix_answers_400_with_the_plan(self):
        self.grant('can_create_blowing_run', 'can_complete_blowing_run')
        resp = self.client.post(self.URL, {
            'machine_id': self.machine.id, 'commit': True,
            'rows': [self._row(date(2026, 9, 30), 'NIGHT', preform_spec_id=None)],
        }, format='json')
        self.assertEqual(resp.status_code, 400)
        self.assertFalse(resp.data['committed'])
        self.assertEqual(resp.data['rows'][0]['status'], 'ERROR')

    def test_bad_requests(self):
        self.grant('can_create_blowing_run', 'can_complete_blowing_run')
        self.assertEqual(self.client.post(self.URL, {'machine_id': self.machine.id, 'rows': []},
                                          format='json').status_code, 400)
        other = Company.objects.create(name='Other', code='OTHER_SS')
        foreign = BlowingMachine.objects.create(company=other, name='Theirs')
        resp = self.client.post(self.URL, self.body(machine_id=foreign.id), format='json')
        self.assertEqual(resp.status_code, 400)
        resp = self.client.post(self.PARSE_URL, {}, format='multipart')
        self.assertEqual(resp.status_code, 400)
