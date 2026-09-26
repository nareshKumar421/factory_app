import io
from datetime import date

from django.test import SimpleTestCase
from openpyxl import Workbook

from tomorrow_run.sheet_parser import SheetError, date_in_name, parse_sheet


def _book(rows, title="PRODUCTION PLANING MONTH OF SEP 2026", tab="P", lead_tab=True):
    wb = Workbook()
    if lead_tab:
        wb.active.title = "SUMMARY"
        wb.active["A1"] = "totals only"
        ws = wb.create_sheet(tab)
    else:
        ws = wb.active
        ws.title = tab
    ws["A1"] = title
    ws.append([])
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


HEADER = ["S.NO", "CODE", "ITEM NAME", "SEP PLAN", "ECOM", "STOCK BH-PF", "STOCK BH-BT", "NET REQ", "MACHINE"]


class ParseTests(SimpleTestCase):
    def test_finds_the_tab_and_the_columns_by_their_headers(self):
        p = parse_sheet(_book([
            HEADER,
            [1, "FG0000030", "MUSTARD KACHI GHANI 1 LTR", 150000, 70000, 100000, 28971, 91029, "JP"],
            [2, "fg0000004", "COLD PRESS 5 LTR", "150,000", 40000, 1830, 0, 188170, "Clearpack 5"],
            [3, None, "PREMIUM OLIVE OIL 500ML", 1000, 0, 0, 0, 1000, ""],
            [None, None, "TOTAL", None, None, None, None, 280199, None],
        ]))
        self.assertEqual((p.tab, p.header_row), ("P", 3))
        self.assertEqual(p.title, "PRODUCTION PLANING MONTH OF SEP 2026")
        self.assertEqual(len(p.lines), 3)
        first, second, third = p.lines
        self.assertEqual((first["code"], first["net_l"], first["stock_l"], first["machine"]),
                         ("FG0000030", 91029.0, 128971.0, "JP"))
        self.assertEqual((second["code"], second["plan_l"]), ("FG0000004", 150000.0))
        self.assertEqual((third["code"], third["name"]), ("", "PREMIUM OLIVE OIL 500ML"))
        self.assertIn("stock", p.columns)
        self.assertEqual(p.net_req_l, 91029 + 188170 + 1000)

    def test_columns_can_move(self):
        p = parse_sheet(_book([
            ["MACHINE", "NET REQ.", "ITEM CODE", "PRODUCT", "PLANNING", "E-COM", "STOCK"],
            ["JP", 500, "FG0000030", "MUSTARD", 400, 200, 100],
        ], lead_tab=False))
        ln = p.lines[0]
        self.assertEqual((ln["code"], ln["net_l"], ln["plan_l"], ln["ecom_l"], ln["stock_l"], ln["machine"]),
                         ("FG0000030", 500.0, 400.0, 200.0, 100.0, "JP"))

    def test_no_planning_tab_is_refused(self):
        with self.assertRaises(SheetError):
            parse_sheet(_book([["CODE", "NAME", "QTY"], ["FG1", "X", 1]]))

    def test_not_an_excel_file(self):
        with self.assertRaises(SheetError):
            parse_sheet(io.BytesIO(b"not a workbook"))

    def test_date_in_the_file_name(self):
        self.assertEqual(date_in_name("PLANNING FOR SEP 2026 19.09.2026 20 se closing tk ki planning - Copy.xlsx"),
                         date(2026, 9, 19))
        self.assertEqual(date_in_name("plan 03-10-2026.xlsx"), date(2026, 10, 3))
        self.assertIsNone(date_in_name("PLANNING FOR SEP 2026.xlsx"))
