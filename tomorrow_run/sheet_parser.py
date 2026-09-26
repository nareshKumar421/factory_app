"""Read the planning team's monthly sheet.

The sheet is theirs, not ours, so nothing here assumes a column order: the tab
is the one whose header row has a CODE column and a Net Req column, and every
other column is found by what its header says. What was matched to what is
returned with the lines and shown when the sheet is put in, so a column that
moved is seen the day it moves, not a month of plans later.

Every quantity on the sheet is litres.
"""

import re
from dataclasses import dataclass, field
from datetime import date
from typing import Any, Dict, List, Optional

from openpyxl import load_workbook

HEADER_SCAN_ROWS = 40


class SheetError(ValueError):
    """The file is not a planning sheet this page can read."""


@dataclass
class ParsedSheet:
    tab: str
    title: str
    header_row: int
    columns: Dict[str, str]
    lines: List[Dict[str, Any]] = field(default_factory=list)

    @property
    def net_req_l(self) -> float:
        return sum(x["net_l"] for x in self.lines)


def _text(v) -> str:
    return " ".join(str(v).split()).strip() if v is not None else ""


def _key(v) -> str:
    return re.sub(r"[^A-Z0-9]+", " ", _text(v).upper()).strip()


def _num(v) -> float:
    if v is None or v == "":
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).replace(",", "").strip()
    if s in ("-", "--", "NA", "N/A"):
        return 0.0
    try:
        return float(s)
    except ValueError:
        return 0.0


def _is_code_header(k: str) -> bool:
    return k == "CODE" or k.endswith(" CODE") or k in ("ITEMCODE", "SAP CODE", "ITEM CODE", "FG CODE")


def _is_net_header(k: str) -> bool:
    return "NET" in k.split() and any(w.startswith("REQ") for w in k.split())


def _find_columns(headers: Dict[int, str]) -> Dict[str, List[int]]:
    cols: Dict[str, List[int]] = {"code": [], "name": [], "plan": [], "ecom": [], "stock": [], "net": [], "machine": []}
    for c, raw in sorted(headers.items()):
        k = _key(raw)
        if not k:
            continue
        words = k.split()
        if _is_code_header(k) and not cols["code"]:
            cols["code"].append(c)
        elif _is_net_header(k):
            cols["net"].append(c)
        elif "ECOM" in k.replace(" ", "") or k.startswith("E COM"):
            cols["ecom"].append(c)
        elif "STOCK" in words or k in ("BH PF", "BH BT", "PF", "BT") or k.startswith("BH PF") or k.startswith("BH BT"):
            cols["stock"].append(c)
        elif "MACHINE" in words or words[0] == "LINE":
            cols["machine"].append(c)
        elif any(w.startswith("PLAN") for w in words) and not cols["plan"]:
            cols["plan"].append(c)
        elif not cols["name"] and any(w in words for w in ("ITEM", "DESCRIPTION", "PRODUCT", "NAME", "SKU", "PARTICULARS")):
            cols["name"].append(c)
    return cols


def _letter(c: int) -> str:
    out = ""
    while c:
        c, r = divmod(c - 1, 26)
        out = chr(65 + r) + out
    return out


def parse_sheet(path_or_file) -> ParsedSheet:
    try:
        wb = load_workbook(path_or_file, data_only=True, read_only=True)
    except Exception as e:  # openpyxl raises a zoo of types for a non-xlsx
        raise SheetError(f"This is not an Excel (.xlsx) file this page can open ({e}).") from e

    seen = []
    try:
        for ws in wb.worksheets:
            rows = list(ws.iter_rows(values_only=True))
            for i, row in enumerate(rows[:HEADER_SCAN_ROWS]):
                keys = [_key(v) for v in row]
                if not any(_is_code_header(k) for k in keys) or not any(_is_net_header(k) for k in keys):
                    if any(keys):
                        seen.append(f"{ws.title}!{i + 1}")
                    continue
                headers = {c + 1: _text(v) for c, v in enumerate(row) if _text(v)}
                # a header split over two rows: fill the blanks from the row below
                if i + 1 < len(rows):
                    for c, v in enumerate(rows[i + 1]):
                        if (c + 1) not in headers and _text(v) and not isinstance(v, (int, float)):
                            headers[c + 1] = _text(v)
                cols = _find_columns(headers)
                if not cols["code"] or not cols["net"]:
                    continue
                return _read(ws.title, rows, i, headers, cols)
    finally:
        wb.close()
    raise SheetError(
        "No tab in this file has a header row with a CODE column and a Net Req column. "
        "The planning sheet's tab has both."
    )


def _read(tab, rows, header_i, headers, cols) -> ParsedSheet:
    title = ""
    for row in rows[:header_i]:
        texts = [_text(v) for v in row if _text(v) and not isinstance(v, (int, float))]
        if texts:
            title = max(texts, key=len)
            break

    def at(row, c):
        return row[c - 1] if c - 1 < len(row) else None

    code_c, net_c = cols["code"][0], cols["net"][0]
    lines = []
    for r_i in range(header_i + 1, len(rows)):
        row = rows[r_i]
        code = _text(at(row, code_c)).upper()
        name = _text(at(row, cols["name"][0])) if cols["name"] else ""
        if _is_code_header(_key(code)):
            continue                     # the second header row
        if not code and not name:
            continue
        if not code and ("TOTAL" in name.upper() or not any(isinstance(at(row, c), (int, float)) for c in (net_c,))):
            continue
        lines.append({
            "row": r_i + 1,
            "code": code,
            "name": name,
            "plan_l": sum(_num(at(row, c)) for c in cols["plan"]),
            "ecom_l": sum(_num(at(row, c)) for c in cols["ecom"]),
            "stock_l": sum(_num(at(row, c)) for c in cols["stock"]),
            "net_l": _num(at(row, net_c)),
            "machine": _text(at(row, cols["machine"][0])) if cols["machine"] else "",
        })

    columns = {
        f: ", ".join(f"{_letter(c)} “{headers.get(c, '')}”" for c in cs)
        for f, cs in cols.items() if cs
    }
    if not lines:
        raise SheetError(f"The tab {tab} has the right header row but no lines under it.")
    return ParsedSheet(tab=tab, title=title, header_row=header_i + 1, columns=columns, lines=lines)


_DATE = re.compile(r"(?<!\d)(\d{1,2})[.\-/_ ](\d{1,2})[.\-/_ ](\d{4})(?!\d)")


def date_in_name(file_name: str) -> Optional[date]:
    """The stock date the planning team writes into the file name ("… 19.09.2026 …")."""
    for m in _DATE.finditer(file_name or ""):
        d, mth, y = (int(x) for x in m.groups())
        try:
            return date(y, mth, d)
        except ValueError:
            continue
    return None
