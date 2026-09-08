"""Excel export of the box scanning done on a Docking load.

The docking review screen answers "did this truck get scanned?" on screen, but the
answer travels: a short bill gets argued about with the warehouse, and a scanned
box gets traced back to its pallet days later. This turns that screen into an
.xlsx the operator can attach to a mail.

Scope follows the review screen exactly. A physical truck can carry more than one
docking -- a cross-company load, or a same-company split (see
``arrival_docking_count``) -- and the screen shows the whole load, so the report
does too: every active docking on the arrival, with a Docking/Company column
appearing only when the load actually has more than one.

Pieces, not boxes, are the unit here. A bill's expected figure is its SAP
quantity; a scan's contribution is the box's own quantity. That is the same
arithmetic the screen's "2,146 / 2,199 pcs" badge does, so the two never
disagree.
"""

from collections import defaultdict
from dataclasses import dataclass
from decimal import Decimal
from typing import Dict, List, Optional, Sequence

from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from gate_core.services.sales_dispatch_gatepass import is_box_scan_optional

XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

# Statuses a docking has to be in to count as part of the load. Mirrors
# ``gate_core.services.arrival_gatepass._ACTIVE_DOCKING_STATUSES`` -- a rejected or
# cancelled sibling is not on the truck and must not dilute the load's totals.
ACTIVE_DOCKING_STATUSES = (
    "DOCKED",
    "PHOTO_ATTACHED",
    "READY_FOR_GATEPASS",
    "GATEPASS_PRINTED",
    "PRINT_COMMITTED",
    "DISPATCHED",
)

# Packaging-material lines are never scanned (they ship as the cartons themselves),
# so a PM-only bill showing zero scans is correct rather than short. Same rule as
# the BST scan gate's ``is_pm_item_code``.
PM_ITEM_PREFIX = "PM"

STATUS_FULL = "FULLY SCANNED"
STATUS_SHORT = "SHORT"
STATUS_NOT_SCANNED = "NOT SCANNED"
STATUS_PM_ITEM = "PM - scanning exempt"
STATUS_PM_DOC = "PM only - scanning exempt"
# Some companies (Jivo Beverages) don't scan boxes at the factory at all, so their
# bills carry no scans by policy. Calling that "NOT SCANNED" next to a five-figure
# shortfall reads as a failure, which is exactly the wrong story.
STATUS_SCAN_OPTIONAL = "Scanning not required (company)"

_THIN = Side(style="thin", color="D0D0D0")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_HEAD_FILL = PatternFill("solid", fgColor="1F3864")
_HEAD_FONT = Font(bold=True, color="FFFFFF", size=10)
_TITLE_FONT = Font(bold=True, size=14, color="1F3864")
_LABEL_FONT = Font(bold=True, size=10)
_FILL_OK = PatternFill("solid", fgColor="E2EFDA")
_FILL_SHORT = PatternFill("solid", fgColor="FCE4D6")
_FILL_NEUTRAL = PatternFill("solid", fgColor="F2F2F2")
_FILL_TOTAL = PatternFill("solid", fgColor="DDEBF7")

_FMT_INT = "#,##0"
_FMT_QTY = "#,##0.###"
_FMT_WEIGHT = "#,##0.000"
_FMT_PCT = "0.0%"
_FMT_DATE = "dd-mmm-yyyy"
_FMT_DATETIME = "dd-mmm-yyyy hh:mm"
_FMT_TIMESTAMP = "dd-mmm-yyyy hh:mm:ss"


# ---------------------------------------------------------------------------
# Gathering
# ---------------------------------------------------------------------------

def load_dockings(entry) -> List:
    """The dockings the report covers: the whole truck, or just this entry.

    Read across companies deliberately -- a cross-company truck's sibling docking
    belongs to another company, and leaving it out would under-report the load the
    screen shows.
    """
    if not entry.arrival_id:
        return [entry]
    siblings = list(
        entry.__class__.objects
        .filter(arrival_id=entry.arrival_id, is_active=True, status__in=ACTIVE_DOCKING_STATUSES)
        .select_related("company", "vehicle", "transporter", "driver")
        .prefetch_related("documents__items", "box_scans__scanned_by", "box_scans__document")
        .order_by("id")
    )
    # A cancelled or rejected docking is not on the truck, so it never appears among
    # its own siblings -- report it alone rather than dropping the entry that was asked for.
    if len(siblings) <= 1 or not any(sibling.pk == entry.pk for sibling in siblings):
        return [entry]
    # Keep the opened entry's own instance so the header block reads from the
    # object the caller already resolved (and its prefetches).
    return [entry if sibling.pk == entry.pk else sibling for sibling in siblings]


def _is_pm(item_code: str) -> bool:
    return (item_code or "").upper().startswith(PM_ITEM_PREFIX)


def _dec(value) -> Decimal:
    return Decimal(value) if value is not None else Decimal(0)


def _num(value) -> Optional[float]:
    return float(value) if value is not None else None


def _local(value):
    """Naive local datetime -- Excel has no timezone, so store the wall clock."""
    return timezone.localtime(value).replace(tzinfo=None) if value else None


def _user_name(user) -> str:
    if not user:
        return ""
    return getattr(user, "full_name", "") or getattr(user, "username", "")


@dataclass
class DocumentRow:
    docking: object
    document: object
    items: List
    bill_pcs: Decimal
    scanned_pcs: Decimal
    box_count: int
    short_pcs: Decimal
    status: str


def _document_rows(dockings: Sequence) -> List[DocumentRow]:
    rows: List[DocumentRow] = []
    for docking in dockings:
        scans_by_document: Dict[Optional[int], List] = defaultdict(list)
        for scan in docking.box_scans.all():
            if scan.is_active:
                scans_by_document[scan.document_id].append(scan)

        for document in docking.documents.all():
            if not document.is_active:
                continue
            items = sorted(
                (item for item in document.items.all() if item.is_active),
                key=lambda item: item.line_num,
            )
            scans = scans_by_document.get(document.id, [])
            bill_pcs = _dec(document.total_quantity)
            scanned_pcs = sum((_dec(scan.quantity) for scan in scans), Decimal(0))
            short_pcs = bill_pcs - scanned_pcs
            if scanned_pcs:
                status = STATUS_FULL if short_pcs <= 0 else STATUS_SHORT
            elif items and all(_is_pm(item.item_code) for item in items):
                status = STATUS_PM_DOC
            elif is_box_scan_optional(docking):
                status = STATUS_SCAN_OPTIONAL
            else:
                status = STATUS_NOT_SCANNED
            rows.append(DocumentRow(
                docking=docking,
                document=document,
                items=items,
                bill_pcs=bill_pcs,
                scanned_pcs=scanned_pcs,
                box_count=len(scans),
                short_pcs=short_pcs,
                status=status,
            ))
    return rows


def _scanned_by_item(dockings: Sequence) -> Dict:
    """Scanned pieces, box count and batches per (document, item code).

    A scan records the item it carries but not the bill *line*, so this is the
    finest grain the data supports.
    """
    totals: Dict = defaultdict(lambda: {"pcs": Decimal(0), "boxes": 0, "batches": set()})
    for docking in dockings:
        for scan in docking.box_scans.all():
            if not scan.is_active:
                continue
            bucket = totals[(scan.document_id, scan.item_code)]
            bucket["pcs"] += _dec(scan.quantity)
            bucket["boxes"] += 1
            if scan.batch_number:
                bucket["batches"].add(scan.batch_number.strip())
    return totals


def _all_scans(dockings: Sequence) -> List:
    """Every active scan on the load, oldest first, paired with its own docking.

    The docking travels alongside rather than through ``scan.sales_dispatch``, which
    would be one query per scan on a load of a few hundred boxes.
    """
    scans = [
        (docking, scan)
        for docking in dockings
        for scan in docking.box_scans.all()
        if scan.is_active
    ]
    return sorted(scans, key=lambda pair: pair[1].scanned_at)


# ---------------------------------------------------------------------------
# Sheet plumbing
# ---------------------------------------------------------------------------

@dataclass(frozen=True)
class Column:
    label: str
    width: float
    number_format: str = ""


def _write_head(ws, columns: Sequence[Column], row: int = 1) -> None:
    for index, column in enumerate(columns, start=1):
        cell = ws.cell(row=row, column=index, value=column.label)
        cell.fill, cell.font, cell.border = _HEAD_FILL, _HEAD_FONT, _BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(index)].width = column.width
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def _write_row(ws, row: int, columns: Sequence[Column], values: Sequence, fill=None) -> None:
    for index, (column, value) in enumerate(zip(columns, values), start=1):
        cell = ws.cell(row=row, column=index, value=value)
        cell.border = _BORDER
        if column.number_format:
            cell.number_format = column.number_format
        if fill is not None:
            cell.fill = fill


def _status_fill(status: str):
    if status == STATUS_FULL:
        return _FILL_OK
    if status in (STATUS_NOT_SCANNED, STATUS_PM_ITEM, STATUS_PM_DOC, STATUS_SCAN_OPTIONAL):
        return _FILL_NEUTRAL
    return _FILL_SHORT


# ---------------------------------------------------------------------------
# Sheets
# ---------------------------------------------------------------------------

def _summary_sheet(ws, entry, dockings, rows: Sequence[DocumentRow], multi: bool) -> None:
    ws.title = "Summary"
    ws["A1"] = "Docking Box-Scan Report"
    ws["A1"].font = _TITLE_FONT
    ws.merge_cells("A1:F1")

    bilty_date = entry.bilty_date.isoformat() if entry.bilty_date else "-"
    meta = [
        ("Docking Entry", entry.entry_no),
        ("Company", entry.company.name),
        ("Status", entry.status),
        ("Vehicle No", entry.vehicle_no or (entry.vehicle.vehicle_number if entry.vehicle_id else "")),
        ("Transporter", entry.transporter_name),
        ("Driver", entry.driver_name),
        ("Bilty No / Date", f"{entry.bilty_no or '-'}  /  {bilty_date}"),
        ("Gatepass No", entry.gatepass_no or "-"),
        ("Docked At", _local(entry.docked_at)),
        ("Dispatched At", _local(entry.dispatched_at)),
        ("Dock Incharge", entry.dock_incharge),
        ("Seal Number", entry.seal_number or "-"),
    ]
    if multi:
        # Say so explicitly: the totals below cover bills this entry does not own.
        meta.append((
            "Load Dockings",
            ", ".join(f"{d.entry_no} ({d.company.name})" for d in dockings),
        ))
    meta.append(("Report Generated", _local(timezone.now())))

    row = 3
    for label, value in meta:
        ws.cell(row=row, column=1, value=label).font = _LABEL_FONT
        cell = ws.cell(row=row, column=2, value=value)
        if label in ("Docked At", "Dispatched At", "Report Generated"):
            cell.number_format = _FMT_DATETIME
        row += 1

    row += 1
    ws.cell(row=row, column=1, value="Scanning by SAP Document").font = _TITLE_FONT
    row += 1

    columns = [
        Column("SAP Doc No", 18),
        Column("Type", 16),
        Column("Doc Date", 13, _FMT_DATE),
        Column("Customer Code", 15),
        Column("Customer Name", 34),
        Column("Bill Pcs", 11, _FMT_INT),
        Column("Scanned Pcs", 12, _FMT_INT),
        Column("Boxes Scanned", 14, _FMT_INT),
        Column("Short Pcs", 11, _FMT_INT),
        Column("Scan %", 9, _FMT_PCT),
        Column("Scan Status", 28),
    ]
    if multi:
        columns = [Column("Docking", 22), Column("Company", 18)] + columns
    _write_head(ws, columns, row=row)

    row += 1
    for doc_row in rows:
        values = [
            doc_row.document.sap_doc_num,
            doc_row.document.document_type,
            doc_row.document.sap_doc_date,
            doc_row.document.customer_code,
            doc_row.document.customer_name,
            _num(doc_row.bill_pcs),
            _num(doc_row.scanned_pcs),
            doc_row.box_count,
            _num(doc_row.short_pcs),
            float(doc_row.scanned_pcs / doc_row.bill_pcs) if doc_row.bill_pcs else None,
            doc_row.status,
        ]
        if multi:
            values = [doc_row.docking.entry_no, doc_row.docking.company.name] + values
        _write_row(ws, row, columns, values, fill=_status_fill(doc_row.status))
        row += 1

    bill_total = sum((r.bill_pcs for r in rows), Decimal(0))
    scanned_total = sum((r.scanned_pcs for r in rows), Decimal(0))
    tail = [
        _num(bill_total), _num(scanned_total), sum(r.box_count for r in rows),
        _num(bill_total - scanned_total),
        float(scanned_total / bill_total) if bill_total else None,
        "",
    ]
    totals = ["TOTAL"] + [""] * (len(columns) - len(tail) - 1) + tail
    _write_row(ws, row, columns, totals, fill=_FILL_TOTAL)
    for index in range(1, len(columns) + 1):
        ws.cell(row=row, column=index).font = Font(bold=True)


def _item_sheet(ws, rows: Sequence[DocumentRow], scanned_items: Dict, multi: bool) -> None:
    columns = [
        Column("SAP Doc No", 16),
        Column("Customer", 28),
        Column("Line", 7),
        Column("Item Code", 14),
        Column("Item Name", 40),
        Column("UOM", 7),
        Column("Bill Qty (Pcs)", 14, _FMT_QTY),
        Column("Scanned Pcs", 12, _FMT_QTY),
        Column("Boxes Scanned", 14),
        Column("Short Pcs", 11, _FMT_QTY),
        Column("Batch(es) Scanned", 26),
        Column("Warehouse", 12),
        Column("Litres", 11, _FMT_WEIGHT),
        Column("Weight (Kg)", 13, _FMT_WEIGHT),
        Column("Scan Status", 22),
    ]
    if multi:
        columns = [Column("Docking", 22)] + columns
    _write_head(ws, columns)

    row = 2
    for doc_row in rows:
        for item in doc_row.items:
            scanned = scanned_items.get((doc_row.document.id, item.item_code))
            # A bill can invoice the same item on two lines (different rate or batch).
            # Scans carry no line number, so the item's scanned figure is apportioned
            # across those lines by each one's share of the item's quantity, and the
            # box count is left blank rather than double-counted.
            same_code = [i for i in doc_row.items if i.item_code == item.item_code]
            if scanned and len(same_code) > 1:
                code_total = sum((_dec(i.quantity) for i in same_code), Decimal(0)) or Decimal(1)
                share = _dec(item.quantity) / code_total
                scanned_pcs = (scanned["pcs"] * share).quantize(Decimal("0.001"))
                box_count = ""
            elif scanned:
                scanned_pcs, box_count = scanned["pcs"], scanned["boxes"]
            else:
                scanned_pcs, box_count = Decimal(0), 0

            bill_qty = _dec(item.quantity)
            short_pcs = bill_qty - scanned_pcs
            if _is_pm(item.item_code):
                status = STATUS_PM_ITEM
            elif not scanned_pcs:
                status = (
                    STATUS_SCAN_OPTIONAL if is_box_scan_optional(doc_row.docking)
                    else STATUS_NOT_SCANNED
                )
            elif short_pcs <= 0:
                status = STATUS_FULL
            else:
                status = STATUS_SHORT

            values = [
                doc_row.document.sap_doc_num,
                doc_row.document.customer_name,
                item.line_num,
                item.item_code,
                item.item_name,
                item.uom,
                _num(bill_qty),
                _num(scanned_pcs),
                box_count,
                _num(short_pcs),
                ", ".join(sorted(scanned["batches"])) if scanned else "",
                item.warehouse_code,
                _num(item.total_litres),
                _num(item.total_weight),
                status,
            ]
            if multi:
                values = [doc_row.docking.entry_no] + values
            _write_row(ws, row, columns, values, fill=_status_fill(status))
            row += 1

    if row > 2:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{row - 1}"


def _scan_sheet(ws, scans: Sequence, multi: bool) -> None:
    columns = [
        Column("#", 5),
        Column("Box Barcode", 24),
        Column("SAP Doc No", 15),
        Column("Customer", 28),
        Column("Item Code", 13),
        Column("Item Name", 38),
        Column("Batch", 14),
        Column("Qty (Pcs)", 10, _FMT_QTY),
        Column("UOM", 7),
        Column("Net Wt", 10, _FMT_WEIGHT),
        Column("Gross Wt", 10, _FMT_WEIGHT),
        Column("Box Status", 12),
        Column("Warehouse", 12),
        Column("Pallet", 24),
        Column("Scanned By", 20),
        Column("Scanned At", 20, _FMT_TIMESTAMP),
    ]
    if multi:
        columns = columns[:1] + [Column("Docking", 22)] + columns[1:]
    _write_head(ws, columns)

    row = 2
    for index, (docking, scan) in enumerate(scans, start=1):
        document = scan.document
        values = [
            index,
            scan.box_barcode,
            document.sap_doc_num if document else "(unattributed)",
            document.customer_name if document else "",
            scan.item_code,
            scan.item_name,
            scan.batch_number,
            _num(scan.quantity),
            scan.uom,
            _num(scan.net_weight),
            _num(scan.gross_weight),
            scan.box_status,
            scan.warehouse_code,
            scan.pallet_code,
            _user_name(scan.scanned_by),
            _local(scan.scanned_at),
        ]
        if multi:
            values = values[:1] + [docking.entry_no] + values[1:]
        _write_row(ws, row, columns, values)
        row += 1

    if row > 2:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{row - 1}"


# ---------------------------------------------------------------------------
# Entry points
# ---------------------------------------------------------------------------

def build_scan_report_workbook(entry, dockings: Optional[Sequence] = None) -> Workbook:
    dockings = list(dockings) if dockings is not None else load_dockings(entry)
    multi = len(dockings) > 1
    rows = _document_rows(dockings)

    workbook = Workbook()
    _summary_sheet(workbook.active, entry, dockings, rows, multi)
    _item_sheet(workbook.create_sheet("Item-wise Scanning"), rows, _scanned_by_item(dockings), multi)
    _scan_sheet(workbook.create_sheet("Box Scans"), _all_scans(dockings), multi)
    return workbook


def build_scan_report_filename(entry) -> str:
    stamp = timezone.localtime().strftime("%Y%m%d-%H%M")
    stem = (entry.entry_no or f"docking-{entry.pk}").replace("/", "-")
    return f"Docking_{stem}_Scan_Report_{stamp}.xlsx"


def scan_report_response(entry, dockings: Optional[Sequence] = None) -> HttpResponse:
    workbook = build_scan_report_workbook(entry, dockings)
    response = HttpResponse(content_type=XLSX_CONTENT_TYPE)
    response["Content-Disposition"] = f'attachment; filename="{build_scan_report_filename(entry)}"'
    # The browser can only read the filename off a cross-origin download when the
    # header is exposed; the frontend names the saved file from it.
    response["Access-Control-Expose-Headers"] = "Content-Disposition"
    workbook.save(response)
    return response
