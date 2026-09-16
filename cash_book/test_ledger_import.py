"""
Tests for reading the card sheet and the person tabs.

The judgement these encode is that **the tab is the person**. A person tab
stacks several blocks, each with its own header, and names other people all
through its detail text -- and none of that changes whose account it is. The
proof is the tab's own Total column, which nets every row in it regardless of
the names and lands on that person's balance. So the importer checks itself
against that figure, and so do these.
"""

from datetime import date, datetime
from io import BytesIO

import openpyxl
from django.test import TestCase

from . import sheet_ledgers
from .sheet_import import SheetError

ATM_HEADER = [
    ["Imprest Received from Vicky Vg"],
    ["Ginni Vg Imprest Debit Card (Vishal)"],
    ["Date", "Op Bal", "Amount Received", "Amount Withdrawl", "Cl Bal"],
]

LEDGER_HEADER = ["Sr. no.", "Date", "Detail", "Out", "In", "Total"]


def sheet_of(rows, title="Sheet1"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = title
    for row in rows:
        ws.append(row)
    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    return openpyxl.load_workbook(stream, data_only=True)[title]


class PersonOfTests(TestCase):
    def test_a_tab_names_its_person(self):
        self.assertEqual(sheet_ledgers.person_of("bunty in out"), "Bunty Ji")
        self.assertEqual(sheet_ledgers.person_of("Jasmeet in out"), "Jasmeet Ji")

    def test_an_unknown_tab_falls_back_to_its_own_name(self):
        self.assertEqual(sheet_ledgers.person_of("Milkha in out"), "Milkha")

    def test_only_in_out_tabs_are_person_ledgers(self):
        wb = openpyxl.Workbook()
        wb.active.title = "Cash details 04-06-2026"
        wb.create_sheet("bunty in out")
        wb.create_sheet("Atm details 04-06-2026")
        wb.create_sheet("Send")
        self.assertEqual(sheet_ledgers.person_sheet_names(wb), ["bunty in out"])


class CardSheetTests(TestCase):
    #: The sheet's own opening rows, with its Cl Bal column.
    ROWS = ATM_HEADER + [
        ["15-05-2026", 19538, None, None, 19538],
        [datetime(2026, 4, 6), None, 100000, None, 119538],
        [datetime(2026, 4, 6), None, None, 50000, 69538],
        [datetime(2026, 6, 6), None, None, 50000, 19538],
    ]

    def test_the_card_is_named_from_above_the_table(self):
        card = sheet_ledgers.read_atm_sheet(sheet_of(self.ROWS))
        self.assertEqual(card["name"], "Ginni Vg Imprest Debit Card (Vishal)")

    def test_the_opening_balance_is_the_first_op_bal(self):
        card = sheet_ledgers.read_atm_sheet(sheet_of(self.ROWS))
        self.assertEqual(card["opening_balance"], 19538)

    def test_receipts_and_withdrawals_are_told_apart_by_column(self):
        card = sheet_ledgers.read_atm_sheet(sheet_of(self.ROWS))
        self.assertEqual(
            [(m["kind"], m["amount"]) for m in card["movements"]],
            [("RECEIPT", 100000), ("WITHDRAWAL", 50000), ("WITHDRAWAL", 50000)],
        )

    def test_a_row_with_neither_amount_is_not_a_movement(self):
        """The opening row carries a balance and nothing else."""
        card = sheet_ledgers.read_atm_sheet(sheet_of(self.ROWS))
        self.assertEqual(len(card["movements"]), 3)

    def test_dates_are_transposed_like_the_rest_of_the_workbook(self):
        card = sheet_ledgers.read_atm_sheet(sheet_of(self.ROWS))
        # Stored Apr 6, typed 04/06 -- 4 June.
        self.assertEqual(card["movements"][0]["date"], date(2026, 6, 4))

    def test_the_arithmetic_reproduces_the_sheets_own_cl_bal(self):
        card = sheet_ledgers.read_atm_sheet(sheet_of(self.ROWS))
        running = card["opening_balance"]
        for movement in card["movements"]:
            running += (
                movement["amount"]
                if movement["kind"] == "RECEIPT"
                else -movement["amount"]
            )
            self.assertEqual(running, movement["stated_balance"])

    def test_the_voucher_log_sharing_the_tab_is_ignored(self):
        rows = [row[:] for row in self.ROWS]
        rows[3] += [datetime(2026, 5, 6), 17570]
        card = sheet_ledgers.read_atm_sheet(sheet_of(rows))
        self.assertEqual(len(card["movements"]), 3)

    def test_a_sheet_with_no_body_is_refused(self):
        with self.assertRaises(SheetError):
            sheet_ledgers.read_atm_sheet(sheet_of(ATM_HEADER))


class PersonSheetTests(TestCase):
    #: Two blocks, a repeated header, and other people named all through --
    #: exactly the shape of "bunty in out".
    ROWS = [
        LEDGER_HEADER,
        [1, datetime(2026, 5, 6), "Bunty jo ko deye", 15000, None, 15000],
        [2, datetime(2026, 10, 6), "Cash muje deya", None, 9000, 6000],
        [3, datetime(2026, 11, 6), "Voucher deye", None, 5000, 1000],
        LEDGER_HEADER,
        [1, datetime(2026, 12, 6), "Tiwari ji ko deye", 10000, None, 11000],
        [None, datetime(2026, 12, 6), "Milkha ne deye", None, 2000, 9000],
    ]

    def ledger(self, rows=None):
        return sheet_ledgers.read_person_sheet(
            sheet_of(rows or self.ROWS, "bunty in out"), "bunty in out"
        )

    def test_every_row_belongs_to_the_tabs_person(self):
        """Tiwari and Milkha are narrative, not a change of account."""
        ledger = self.ledger()
        self.assertEqual({row["person"] for row in ledger["rows"]}, {"Bunty Ji"})

    def test_a_repeated_header_is_skipped_not_treated_as_a_break(self):
        ledger = self.ledger()
        self.assertEqual(len(ledger["rows"]), 5)

    def test_the_out_column_is_cash_given_and_the_in_column_clears_it(self):
        ledger = self.ledger()
        self.assertEqual(
            [(row["direction"], row["amount"]) for row in ledger["rows"]],
            [
                ("GIVEN", 15000),
                ("CLEARED", 9000),
                ("CLEARED", 5000),
                ("GIVEN", 10000),
                ("CLEARED", 2000),
            ],
        )

    def test_the_arithmetic_lands_on_the_tabs_own_total(self):
        """The check the importer makes, and the reason the tab is the person."""
        ledger = self.ledger()
        net = sum(
            row["amount"] if row["direction"] == "GIVEN" else -row["amount"]
            for row in ledger["rows"]
        )
        self.assertEqual(net, ledger["stated_balance"])
        self.assertEqual(net, 9000)

    def test_a_voucher_row_is_recognised_however_it_is_spelled(self):
        rows = [
            LEDGER_HEADER,
            [1, datetime(2026, 5, 6), "Voucher deye", None, 100, -100],
            [2, datetime(2026, 5, 6), "Vocher deye", None, 100, -200],
            [3, datetime(2026, 5, 6), "Vouher deye", None, 100, -300],
            [4, datetime(2026, 5, 6), "Cash muje deya", None, 100, -400],
        ]
        ledger = self.ledger(rows)
        self.assertEqual(
            [row["is_voucher"] for row in ledger["rows"]], [True, True, True, False]
        )

    def test_an_undated_row_is_kept_and_left_for_the_importer_to_date(self):
        rows = [
            LEDGER_HEADER,
            [1, None, "Tiwari ji ko deye", 10000, None, 10000],
        ]
        ledger = self.ledger(rows)
        self.assertEqual(len(ledger["rows"]), 1)
        self.assertIsNone(ledger["rows"][0]["date"])

    def test_a_row_with_no_amount_either_side_is_not_a_movement(self):
        rows = [
            LEDGER_HEADER,
            [1, datetime(2026, 5, 6), "Bunty jo ko deye", 15000, None, 15000],
            [2, None, None, None, None, 15000],
        ]
        self.assertEqual(len(self.ledger(rows)["rows"]), 1)
