"""
Connect each app login to the employee it belongs to.

``Employee.user`` is the only thing that answers "which person is signed in?",
and until it is set the answer is *nobody*. Measured on the live database when
this command was written: **249 employees, 128 logins, 0 links** -- and no way
to derive one, because no employee carries an email and no login carries a
matching ``employee_code``. The two sets were built by different people at
different times and never introduced.

Everything that is about *you* needs this link. Your own salary
(:mod:`employee_hierarchy.access` refuses to guess which employee an unlinked
login is, deliberately -- guessing would be guessing about somebody's pay), and
now leave: who is applying, and whose manager should decide it.

**Three ways to match, in descending order of trust.** Each is opt-in, because
a wrong link is not a visible error -- it is one person applying for leave as
another, and later one person reading another's salary.

``--file``
    A sheet HR fills in: ``employee_code`` and ``email`` columns. The only
    mode that can cover somebody whose login shares nothing with their
    employee record, which on this data is everybody. ``.csv`` or ``.xlsx``.

``--by-employee-code``
    Match ``accounts_user.employee_code`` to ``Employee.employee_code``.
    Exact, and free when it works -- it matched 0 people on the live data,
    but new users created properly will carry it.

``--by-email``
    Match ``accounts_user.email`` to ``Employee.email``, case-insensitively.
    Only ever acts where the address appears **once** on each side; an address
    two employees share is reported, never guessed.

Dry run by default, like the other repairs in this module::

    # look first -- writes nothing
    python manage.py link_employee_logins --file hr_logins.xlsx

    # do it
    python manage.py link_employee_logins --file hr_logins.xlsx --commit

    # the free ones, together
    python manage.py link_employee_logins --by-employee-code --by-email --commit

Safe to re-run. A link that already points where it should is left alone and
counted as "already linked"; one that points somewhere *else* is refused and
reported, never silently moved, because a login changing hands is an HR event
and not a side effect of re-running an import.
"""

import csv
from collections import Counter, defaultdict
from pathlib import Path

from django.contrib.auth import get_user_model
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from company.models import Company

from ...models import Employee

User = get_user_model()

#: Code cells that mean "no code", not a code. Same vocabulary as
#: ``import_employee_branches`` -- the sheets come from the same place.
NOT_A_CODE = {"", "NEW SEP", "NONE", "NA", "N/A", "-"}


class Command(BaseCommand):
    help = "Link app logins to their employee records (dry run unless --commit)."

    def add_arguments(self, parser):
        parser.add_argument(
            "--file",
            help="Sheet with employee_code and email columns (.csv or .xlsx).",
        )
        parser.add_argument(
            "--by-employee-code",
            action="store_true",
            help="Also match accounts_user.employee_code to Employee.employee_code.",
        )
        parser.add_argument(
            "--by-email",
            action="store_true",
            help="Also match accounts_user.email to Employee.email (unique on both sides only).",
        )
        parser.add_argument(
            "--company",
            help="Restrict to one company code, e.g. JIVO_OIL. Default: every company.",
        )
        parser.add_argument(
            "--commit",
            action="store_true",
            help="Actually write. Without this nothing is saved.",
        )

    def handle(self, *args, **options):
        if not any(
            (options.get("file"), options["by_employee_code"], options["by_email"])
        ):
            raise CommandError(
                "Nothing to do. Pass --file, --by-employee-code or --by-email "
                "(the last two can be combined)."
            )

        company = None
        if options.get("company"):
            company = Company.objects.filter(code=options["company"]).first()
            if company is None:
                raise CommandError(f"No company with code {options['company']!r}.")

        employees = Employee.objects.all()
        if company is not None:
            employees = employees.filter(company=company)

        #: proposals maps user_id -> (employee, how it was matched). One login
        #: can only be one employee, so a second proposal for the same login is
        #: a conflict rather than an overwrite.
        proposals = {}
        conflicts = []
        notes = Counter()

        if options.get("file"):
            self._match_from_file(
                Path(options["file"]), employees, proposals, conflicts, notes
            )
        if options["by_employee_code"]:
            self._match_by_employee_code(employees, proposals, conflicts, notes)
        if options["by_email"]:
            self._match_by_email(employees, proposals, conflicts, notes)

        self._report_and_write(proposals, conflicts, notes, commit=options["commit"])

    # -- the three matchers --------------------------------------------------

    def _match_from_file(self, path, employees, proposals, conflicts, notes):
        if not path.exists():
            raise CommandError(f"No such file: {path}")

        rows = self._read_rows(path)
        by_code = {
            (employee.company_id, employee.employee_code.upper()): employee
            for employee in employees
        }
        users_by_email = {}
        for user in User.objects.all():
            users_by_email.setdefault(user.email.strip().lower(), []).append(user)

        for line_no, row in rows:
            code = (row.get("employee_code") or "").strip().upper()
            email = (row.get("email") or "").strip().lower()
            if code in NOT_A_CODE or not email:
                notes["file rows skipped (blank code or email)"] += 1
                continue

            matched = [
                employee
                for (_, employee_code), employee in by_code.items()
                if employee_code == code
            ]
            if not matched:
                conflicts.append(f"row {line_no}: no employee with code {code}")
                continue
            if len(matched) > 1:
                conflicts.append(
                    f"row {line_no}: code {code} exists in "
                    f"{len(matched)} companies -- pass --company to disambiguate"
                )
                continue

            candidates = users_by_email.get(email, [])
            if not candidates:
                conflicts.append(f"row {line_no}: no login with email {email}")
                continue
            if len(candidates) > 1:
                conflicts.append(f"row {line_no}: {len(candidates)} logins share {email}")
                continue

            self._propose(proposals, conflicts, candidates[0], matched[0], "file")

    def _match_by_employee_code(self, employees, proposals, conflicts, notes):
        by_code = defaultdict(list)
        for employee in employees:
            by_code[employee.employee_code.strip().upper()].append(employee)

        for user in User.objects.exclude(employee_code__isnull=True).exclude(
            employee_code=""
        ):
            code = user.employee_code.strip().upper()
            if code in NOT_A_CODE:
                continue
            matched = by_code.get(code, [])
            if not matched:
                notes["logins whose employee_code matches nobody"] += 1
                continue
            if len(matched) > 1:
                conflicts.append(
                    f"login {user.email}: code {code} exists in {len(matched)} companies"
                )
                continue
            self._propose(proposals, conflicts, user, matched[0], "employee_code")

    def _match_by_email(self, employees, proposals, conflicts, notes):
        # Only addresses that appear exactly once on each side. An address two
        # people share identifies neither of them.
        employees_by_email = defaultdict(list)
        for employee in employees:
            address = (employee.email or "").strip().lower()
            if address:
                employees_by_email[address].append(employee)

        users_by_email = defaultdict(list)
        for user in User.objects.all():
            address = (user.email or "").strip().lower()
            if address:
                users_by_email[address].append(user)

        for address, matched in employees_by_email.items():
            candidates = users_by_email.get(address, [])
            if not candidates:
                continue
            if len(matched) > 1 or len(candidates) > 1:
                conflicts.append(
                    f"{address}: {len(matched)} employee(s) and {len(candidates)} login(s) "
                    "-- ambiguous, skipped"
                )
                continue
            self._propose(proposals, conflicts, candidates[0], matched[0], "email")

    # -- shared ---------------------------------------------------------------

    def _propose(self, proposals, conflicts, user, employee, how):
        """Record that ``user`` should be ``employee``, or explain why not."""
        existing = proposals.get(user.pk)
        if existing is not None and existing[0].pk != employee.pk:
            conflicts.append(
                f"login {user.email}: matched to both {existing[0].employee_code} "
                f"({existing[1]}) and {employee.employee_code} ({how}) -- skipped"
            )
            proposals.pop(user.pk, None)
            return
        proposals[user.pk] = (employee, how)

    def _read_rows(self, path):
        """``[(line_no, {column: value})]`` from a .csv or .xlsx."""
        if path.suffix.lower() == ".csv":
            with path.open(newline="", encoding="utf-8-sig") as handle:
                return [
                    (index, {(k or "").strip().lower(): v for k, v in row.items()})
                    for index, row in enumerate(csv.DictReader(handle), start=2)
                ]

        try:
            from openpyxl import load_workbook
        except ImportError as exc:  # pragma: no cover - openpyxl is in requirements
            raise CommandError("openpyxl is needed to read .xlsx files.") from exc

        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook.active
        rows = []
        headers = []
        for index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            if index == 1:
                headers = [str(cell or "").strip().lower() for cell in row]
                continue
            rows.append(
                (index, {headers[i]: row[i] for i in range(min(len(headers), len(row)))})
            )
        workbook.close()
        return rows

    def _report_and_write(self, proposals, conflicts, notes, *, commit):
        already = 0
        taken = []
        to_write = []

        for user_id, (employee, how) in sorted(proposals.items()):
            if employee.user_id == user_id:
                already += 1
                continue
            if employee.user_id is not None:
                taken.append(
                    f"{employee.employee_code}: already linked to another login -- left alone"
                )
                continue
            # The other side of the one-to-one: this login may already be
            # somebody else's employee record.
            other = Employee.objects.filter(user_id=user_id).exclude(pk=employee.pk).first()
            if other is not None:
                taken.append(
                    f"{employee.employee_code}: that login is already "
                    f"{other.employee_code} -- left alone"
                )
                continue
            to_write.append((user_id, employee, how))

        self.stdout.write("")
        self.stdout.write(f"  proposed links      : {len(proposals)}")
        self.stdout.write(f"  already correct     : {already}")
        self.stdout.write(f"  to link             : {len(to_write)}")
        self.stdout.write(f"  refused (taken)     : {len(taken)}")
        self.stdout.write(f"  unmatched / ambiguous: {len(conflicts)}")
        for label, count in sorted(notes.items()):
            self.stdout.write(f"  {label}: {count}")

        for line in taken[:20]:
            self.stdout.write(self.style.WARNING(f"    ! {line}"))
        for line in conflicts[:20]:
            self.stdout.write(self.style.WARNING(f"    ? {line}"))
        if len(conflicts) > 20:
            self.stdout.write(f"    ... and {len(conflicts) - 20} more")

        if not commit:
            self.stdout.write("")
            self.stdout.write(
                self.style.WARNING("Dry run -- nothing written. Re-run with --commit.")
            )
            return

        with transaction.atomic():
            for user_id, employee, _how in to_write:
                employee.user_id = user_id
                employee.save(update_fields=["user", "updated_at"])

        self.stdout.write("")
        self.stdout.write(self.style.SUCCESS(f"Linked {len(to_write)} login(s)."))
